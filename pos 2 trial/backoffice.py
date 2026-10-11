import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from tkcalendar import DateEntry
import requests
import datetime
import threading
import os
import csv
import subprocess
import sys
import ctypes

SERVER = "http://localhost:5000"

APP_DIR = os.path.dirname(os.path.abspath(__file__))

ICON_MAX_SIZE = 20
ICON_PADDING = 4
ICON_CANVAS_SIZE = ICON_MAX_SIZE + (ICON_PADDING * 2)

TITLEBAR_COLOR_RGB = (0xA9, 0xC9, 0xA4)

SHOP_NAME = "Evans Shop"


# ---------------------------------------------------------------------------
# Windows helpers
# ---------------------------------------------------------------------------

def apply_green_titlebar(win):
    if not sys.platform.startswith("win"):
        return
    try:
        win.update_idletasks()
        hwnd = ctypes.windll.user32.GetParent(win.winfo_id())
        color = ctypes.c_int(
            (TITLEBAR_COLOR_RGB[2] << 16) | (TITLEBAR_COLOR_RGB[1] << 8) | TITLEBAR_COLOR_RGB[0]
        )
        res = ctypes.windll.dwmapi.DwmSetWindowAttribute(
            hwnd, 35, ctypes.byref(color), ctypes.sizeof(color)
        )
        if res != 0:
            dark = ctypes.c_int(0)
            ctypes.windll.dwmapi.DwmSetWindowAttribute(
                hwnd, 20, ctypes.byref(dark), ctypes.sizeof(dark)
            )
    except Exception:
        pass


def open_with_default_app(path):
    try:
        if sys.platform.startswith("win"):
            os.startfile(path)
        elif sys.platform == "darwin":
            subprocess.Popen(["open", path])
        else:
            subprocess.Popen(["xdg-open", path])
    except Exception as e:
        messagebox.showerror("Could not open file",
                             f"Saved to:\n{path}\n\nBut could not open it automatically:\n{e}")


# ---------------------------------------------------------------------------
# Code128 encoder
# ---------------------------------------------------------------------------

_CODE128_PATTERNS = [
    "11011001100","11001101100","11001100110","10010011000","10010001100",
    "10001001100","10011001000","10011000100","10001100100","11001001000",
    "11001000100","11000100100","10110011100","10011011100","10011001110",
    "10111001100","10011101100","10011100110","11001110010","11001011100",
    "11001001110","11011100100","11001110100","11101101110","11101001100",
    "11100101100","11100100110","11101100100","11100110100","11100110010",
    "11011011000","11011000110","11000110110","10100011000","10001011000",
    "10001000110","10110001000","10001101000","10001100010","11010001000",
    "11000101000","11000100010","10110111000","10110001110","10001101110",
    "10111011000","10111000110","10001110110","11101110110","11010001110",
    "11000101110","11011101000","11011100010","11011101110","11101011000",
    "11101000110","11100010110","11101101000","11101100010","11100011010",
    "11101111010","11001000010","11110001010","10100110000","10100001100",
    "10010110000","10010000110","10000101100","10000100110","10110010000",
    "10110000100","10011010000","10011000010","10000110100","10000110010",
    "11000010010","11001010000","11110111010","11000010100","10001111010",
    "10100111100","10010111100","10010011110","10111100100","10011110100",
    "10011110010","11110100100","11110010100","11110010010","11011011110",
    "11011110110","11110110110","10101111000","10100011110","10001011110",
    "10111101000","10111100010","11110101000","11110100010","10111011110",
    "10111101110","11101011110","11110101110","11010000100","11010010000",
    "11010011100","1100011101011",
]


def code128_encode(data):
    data = str(data) or "0"
    bits = "11010010000"
    checksum = 104
    for i, ch in enumerate(data):
        val = ord(ch) - 32
        if val < 0 or val > 95:
            val = 0
        bits += _CODE128_PATTERNS[val]
        checksum += val * (i + 1)
    checksum %= 103
    bits += _CODE128_PATTERNS[checksum]
    bits += "1100011101011"
    return bits


# ---------------------------------------------------------------------------
# Icon helpers
# ---------------------------------------------------------------------------

def load_icon(filename, max_size=ICON_MAX_SIZE):
    path = os.path.join(APP_DIR, filename)
    try:
        img = tk.PhotoImage(file=path)
        longest_side = max(img.width(), img.height())
        if longest_side > max_size:
            factor = -(-longest_side // max_size)
            img = img.subsample(factor, factor)
        return img
    except Exception as e:
        print(f"Could not load icon '{filename}':", e)
        return None


def make_button_icon(icon, canvas_size=ICON_CANVAS_SIZE):
    if icon is None:
        return None
    try:
        canvas = tk.PhotoImage(width=canvas_size, height=canvas_size)
        canvas.put("white", to=(0, 0, canvas_size, canvas_size))
        icon_w, icon_h = icon.width(), icon.height()
        x = max(0, (canvas_size - icon_w) // 2)
        y = max(0, (canvas_size - icon_h) // 2)
        canvas.tk.call(canvas, "copy", icon, "-to", x, y, x + icon_w, y + icon_h,
                        "-compositingrule", "overlay")
        return canvas
    except Exception as e:
        print("Could not build padded button icon:", e)
        return icon


def center_window_on_screen(win, width, height):
    win.update_idletasks()
    try:
        screen_w = win.winfo_screenwidth()
        screen_h = win.winfo_screenheight()
        x = max(0, (screen_w - width) // 2)
        y = max(0, (screen_h - height) // 2)
        win.geometry(f"{width}x{height}+{x}+{y}")
    except tk.TclError:
        pass


class BackOffice(tk.Tk):
    RIBBON_BG = "#a9c9a4"

    def __init__(self):
        super().__init__()
        self.title("Back Office - Electronics Shop")
        self.geometry("950x650")
        self.minsize(800, 500)
        self.configure(bg=self.RIBBON_BG)

        self.last_sale_id_seen = 0
        self.products_cache = []
        self.open_windows = {}
        self.window_groups = {}
        self.current_group = None
        self.admin_credentials = None
        self.known_departments = set()

        self.label_mode_var = None
        self.label_mode_buttons = {}
        self.label_printer_name = None
        self._label_code_normalizing = False
        self._label_bottom_built = False
        self._label_current_area = None
        self._label_win = None
        self._label_body = None

        self._stock_filter_after_id = None
        self._product_filter_after_id = None

        self._grn_code_normalizing = False
        self._stock_code_normalizing = False
        self._product_code_normalizing = False

        self._reclamping = False
        self._main_configuring = False

        self.icon_stock = load_icon("stock.png")
        self.icon_products = load_icon("products.png")
        self.icon_labels = load_icon("label.png")
        self.icon_purchase = load_icon("purchase.png")
        self.icon_orders = load_icon("orders.png")
        self.icon_inventory = load_icon("inventory.png")
        self.icon_pos = load_icon("pos.png")
        self.icon_shelfs = load_icon("shelfs.png")
        self.icon_grn = load_icon("grn.png")
        self.icon_crn = load_icon("crn.png")

        self.button_icon_stock = make_button_icon(self.icon_stock)
        self.button_icon_products = make_button_icon(self.icon_products)
        self.button_icon_labels = make_button_icon(self.icon_labels)
        self.button_icon_purchase = make_button_icon(self.icon_purchase)
        self.button_icon_orders = make_button_icon(self.icon_orders)
        self.button_icon_inventory = make_button_icon(self.icon_inventory)
        self.button_icon_pos = make_button_icon(self.icon_pos)
        self.button_icon_shelfs = make_button_icon(self.icon_shelfs)
        self.button_icon_grn = make_button_icon(self.icon_grn)
        self.button_icon_crn = make_button_icon(self.icon_crn)

        self.setup_dark_style()

        top_ribbon = ttk.Frame(self, style="Ribbon.TFrame")
        top_ribbon.pack(fill="x", padx=5, pady=5)

        ttk.Button(top_ribbon, text="Inventory", image=self.button_icon_inventory, compound="top",
                   command=self.show_inventory_ribbon).pack(side="left", padx=3)
        ttk.Button(top_ribbon, text="Purchase", image=self.button_icon_purchase, compound="top",
                   command=self.show_purchase_ribbon).pack(side="left", padx=3)
        ttk.Button(top_ribbon, text="Orders", image=self.button_icon_orders, compound="top",
                   command=self.show_orders_ribbon).pack(side="left", padx=3)
        ttk.Button(top_ribbon, text="POS", image=self.button_icon_pos, compound="top",
                   command=self.show_pos_ribbon).pack(side="left", padx=3)

        ttk.Separator(self, orient="horizontal").pack(fill="x")

        self.sub_ribbon = ttk.Frame(self, style="Ribbon.TFrame")
        self.sub_ribbon.pack(fill="x", padx=5, pady=5)

        ttk.Separator(self, orient="horizontal").pack(fill="x")

        self.refresh_products_cache()
        self.periodic_cache_refresh()

        self.after(50, self.maximize_main_window)
        self.after(200, lambda: apply_green_titlebar(self))
        self.bind("<Configure>", self.on_main_configure, add="+")
        self.bind("<FocusIn>", self.raise_all_subwindows, add="+")

    def raise_all_subwindows(self, event=None):
        if event is not None and event.widget != self:
            return
        for key, win in list(self.open_windows.items()):
            if win.winfo_exists() and self.window_groups.get(key) == self.current_group:
                win.lift()

    def switch_group(self, group):
        self.current_group = group
        for key, win in list(self.open_windows.items()):
            if not win.winfo_exists():
                continue
            if self.window_groups.get(key) == group:
                win.deiconify()
                win.lift()
            else:
                win.withdraw()

    def setup_dark_style(self):
        style = ttk.Style()
        style.theme_use("clam")
        style.configure("Dark.Treeview", background="black", foreground="white",
                        fieldbackground="black", bordercolor="black")
        style.configure("Dark.Treeview.Heading", background="#222222", foreground="white")
        style.map("Dark.Treeview", background=[("selected", "#333333")])
        style.configure("Ribbon.TFrame", background=self.RIBBON_BG)

    def maximize_main_window(self):
        try:
            self.state("zoomed")
        except tk.TclError:
            try:
                w = self.winfo_screenwidth()
                h = self.winfo_screenheight()
                self.geometry(f"{w}x{h}+0+0")
            except tk.TclError:
                pass

    def on_main_configure(self, event=None):
        if event is not None and event.widget != self:
            return
        if self._main_configuring:
            return
        self._main_configuring = True
        try:
            for win in list(self.open_windows.values()):
                if win.winfo_exists():
                    self.reclamp_window(win)
        finally:
            self._main_configuring = False

    def make_columns_responsive(self, win, tree, weights):
        columns = tree["columns"]
        total_weight = sum(weights)

        def resize(event=None):
            try:
                width = tree.winfo_width()
                if width <= 1:
                    return
                for col, w in zip(columns, weights):
                    tree.column(col, width=int(width * w / total_weight))
            except tk.TclError:
                pass

        win.bind("<Configure>", resize, add="+")
        win.after(100, resize)

    def reclamp_window(self, win):
        if self._reclamping:
            return
        try:
            main_x = self.winfo_x()
            main_y = self.winfo_y()
            main_w = self.winfo_width()
            main_h = self.winfo_height()
            min_y = self.sub_ribbon.winfo_rooty() + self.sub_ribbon.winfo_height() + 8
            max_x = main_x + main_w
            max_y = main_y + main_h

            x = win.winfo_x()
            y = win.winfo_y()
            w = win.winfo_width()
            h = win.winfo_height()

            new_w = min(w, main_w)
            new_h = min(h, max(150, max_y - min_y))
            new_x = x
            new_y = y

            if new_x < main_x:
                new_x = main_x
            if new_x + new_w > max_x:
                new_x = max_x - new_w
            if new_y < min_y:
                new_y = min_y
            if new_y + new_h > max_y:
                new_y = max_y - new_h
            if new_y < min_y:
                new_y = min_y

            if (new_x, new_y, new_w, new_h) != (x, y, w, h):
                self._reclamping = True
                try:
                    win.geometry(f"{new_w}x{new_h}+{new_x}+{new_y}")
                    win.update_idletasks()
                finally:
                    self._reclamping = False
        except tk.TclError:
            pass

    def constrain_window_to_main(self, win):
        def clamp(event=None):
            if event is not None and event.widget != win:
                return
            self.reclamp_window(win)
        win.bind("<Configure>", clamp, add="+")

    def refresh_products_cache(self):
        def worker():
            try:
                r = requests.get(f"{SERVER}/products", timeout=3)
                data = r.json()
            except Exception as e:
                print("Could not refresh product cache:", e)
                return
            self.after(0, lambda: self._apply_products_cache(data))
        threading.Thread(target=worker, daemon=True).start()

    def _apply_products_cache(self, data):
        self.products_cache = data

    def periodic_cache_refresh(self):
        self.refresh_products_cache()
        self.after(5000, self.periodic_cache_refresh)

    def find_product_by_code(self, code):
        if not code:
            return None
        code = str(code).strip()
        return next(
            (p for p in self.products_cache
             if str(p.get("barcode", "")).strip() == code
             or str(p.get("code", "")).strip() == code),
            None
        )

    def get_short_code(self, product):
        short = product.get("code")
        if short:
            return str(short).strip()
        return str(product.get("barcode", "")).strip()

    def clear_sub_ribbon(self):
        for widget in self.sub_ribbon.winfo_children():
            widget.destroy()

    SUBWINDOW_WIDTH = 750
    SUBWINDOW_HEIGHT = 450

    def open_content_window(self, key, title, group=None, icon=None):
        self.window_groups[key] = group
        existing = self.open_windows.get(key)
        if existing is not None and existing.winfo_exists():
            existing.deiconify()
            existing.lift()
            existing.focus_force()
            return None
        try:
            if self.state() == "iconic":
                self.deiconify()
                self.lift()
                self.focus_force()
                self.update_idletasks()
        except tk.TclError:
            pass
        self.update_idletasks()
        main_w = self.winfo_width()
        main_h = self.winfo_height()
        width = min(self.SUBWINDOW_WIDTH, main_w)
        height = min(self.SUBWINDOW_HEIGHT, max(300, main_h - 120))
        win = tk.Toplevel(self)
        win.title(title)
        win.resizable(True, True)
        win.minsize(700, 400)
        center_window_on_screen(win, width, height)
        win.transient(self)
        if icon is not None:
            try:
                win.iconphoto(False, icon)
            except Exception:
                pass
        self.constrain_window_to_main(win)

        def on_close():
            if key in self.open_windows:
                del self.open_windows[key]
            self.window_groups.pop(key, None)
            win.destroy()

        win.protocol("WM_DELETE_WINDOW", on_close)
        self.open_windows[key] = win
        win.after(150, lambda w=win: apply_green_titlebar(w))
        return win

    # ---------- Report window ----------

    def _open_report_window(self, title, rows, columns, weights, scope_label):
        win = tk.Toplevel(self)
        win.title(title)
        win.transient(self)
        screen_w = win.winfo_screenwidth()
        screen_h = win.winfo_screenheight()
        width = max(900, screen_w - 80)
        height = max(600, screen_h - 120)
        x = (screen_w - width) // 2
        y = (screen_h - height) // 2
        win.geometry(f"{width}x{height}+{x}+{y}")
        win.minsize(800, 500)
        win.lift()
        try:
            win.attributes("-topmost", True)
            win.after(300, lambda w=win: self._unset_topmost(w))
        except tk.TclError:
            pass
        win.focus_force()
        win.after(150, lambda w=win: apply_green_titlebar(w))

        top = ttk.Frame(win)
        top.pack(fill="x", padx=12, pady=(12, 0))
        ttk.Label(top, text=scope_label, font=("Arial", 11, "bold")).pack(side="left")
        ttk.Button(top, text="Export…",
                   command=lambda: self._export_report(win, columns, rows, title)).pack(side="right")
        ttk.Button(top, text="Print",
                   command=lambda: self._print_report(win, columns, rows, title)).pack(side="right", padx=(0, 8))

        table_frame = ttk.Frame(win)
        table_frame.pack(fill="both", expand=True, padx=12, pady=12)
        col_keys = [c[0] for c in columns]
        tree = ttk.Treeview(table_frame, columns=col_keys, show="headings",
                            style="Dark.Treeview")
        for key, heading in columns:
            tree.heading(key, text=heading)
            tree.column(key, anchor="w")
        vsb = ttk.Scrollbar(table_frame, orient="vertical", command=tree.yview)
        hsb = ttk.Scrollbar(table_frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        table_frame.rowconfigure(0, weight=1)
        table_frame.columnconfigure(0, weight=1)
        for row in rows:
            tree.insert("", "end", values=row)
        self.make_columns_responsive(win, tree, weights)
        footer = ttk.Frame(win)
        footer.pack(fill="x", padx=12, pady=(0, 12))
        ttk.Label(footer, text=f"Total rows: {len(rows)}").pack(side="left")
        return win

    def _unset_topmost(self, win):
        try:
            if win.winfo_exists():
                win.attributes("-topmost", False)
        except tk.TclError:
            pass

    # ---------- Export ----------

    def _export_report(self, parent, columns, rows, default_name):
        fmt = {"value": None}
        picker = tk.Toplevel(parent)
        picker.title("Export As…")
        picker.resizable(False, False)
        picker.transient(parent)
        picker.grab_set()
        center_window_on_screen(picker, 320, 160)
        picker.after(150, lambda w=picker: apply_green_titlebar(w))

        ttk.Label(picker, text="Choose export format:").pack(anchor="w", padx=14, pady=(14, 6))
        fmt_var = tk.StringVar(value="Excel Workbook (*.xlsx)")
        ttk.Combobox(picker, textvariable=fmt_var, state="readonly", width=30,
                     values=["Excel Workbook (*.xlsx)", "CSV File (*.csv)"]).pack(padx=14)

        def do_ok():
            fmt["value"] = "xlsx" if fmt_var.get().startswith("Excel") else "csv"
            picker.destroy()

        btn = ttk.Frame(picker)
        btn.pack(fill="x", padx=14, pady=14)
        ttk.Button(btn, text="Cancel", command=picker.destroy).pack(side="right")
        ttk.Button(btn, text="Continue", command=do_ok).pack(side="right", padx=(0, 8))
        picker.wait_window()

        if fmt["value"] is None:
            return
        if fmt["value"] == "xlsx":
            path = filedialog.asksaveasfilename(
                parent=parent, title="Export as Excel workbook",
                defaultextension=".xlsx",
                initialfile=f"{default_name}.xlsx",
                filetypes=[("Excel Workbook", "*.xlsx"), ("All files", "*.*")])
            if not path:
                return
            self._write_xlsx(path, columns, rows)
        else:
            path = filedialog.asksaveasfilename(
                parent=parent, title="Export as CSV",
                defaultextension=".csv",
                initialfile=f"{default_name}.csv",
                filetypes=[("CSV File", "*.csv"), ("All files", "*.*")])
            if not path:
                return
            self._write_csv(path, columns, rows)
        messagebox.showinfo("Export complete", f"Saved to:\n{path}", parent=parent)
        open_with_default_app(path)

    def _write_csv(self, path, columns, rows):
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow([c[1] for c in columns])
            for row in rows:
                writer.writerow(row)

    def _write_xlsx(self, path, columns, rows):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            messagebox.showerror("Missing library",
                                 "Excel export needs the 'openpyxl' package.\n\n"
                                 "Install it with:   pip install openpyxl")
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "Report"
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="222222")
        for col_idx, (_, heading) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=heading)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="left")
        for r_idx, row in enumerate(rows, start=2):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)
        for c_idx, (_, heading) in enumerate(columns, start=1):
            max_len = len(str(heading))
            for row in rows:
                if c_idx - 1 < len(row):
                    max_len = max(max_len, len(str(row[c_idx - 1])))
            ws.column_dimensions[
                ws.cell(row=1, column=c_idx).column_letter
            ].width = min(max_len + 2, 50)
        wb.save(path)

    def _print_report(self, parent, columns, rows, title):
        lines = [title, "=" * len(title), ""]
        headings = [c[1] for c in columns]
        widths = []
        for i, h in enumerate(headings):
            max_w = len(h)
            for row in rows:
                if i < len(row):
                    max_w = max(max_w, len(str(row[i])))
            widths.append(min(max_w, 40))

        def fmt(row):
            cells = []
            for i, cell in enumerate(row):
                s = str(cell)
                if len(s) > widths[i]:
                    s = s[: widths[i] - 1] + "…"
                cells.append(s.ljust(widths[i]))
            return " | ".join(cells)

        lines.append(fmt(headings))
        lines.append("-+-".join("-" * w for w in widths))
        for row in rows:
            lines.append(fmt(row))
        lines.append("")
        lines.append(f"Total rows: {len(rows)}")
        text = "\n".join(lines)
        import tempfile
        safe = "".join(ch for ch in title if ch.isalnum() or ch in " -_")[:60] or "report"
        try:
            tmp_dir = tempfile.gettempdir()
            path = os.path.join(tmp_dir, f"{safe}.txt")
            with open(path, "w", encoding="utf-8") as f:
                f.write(text)
        except Exception as e:
            messagebox.showerror("Print failed",
                                 f"Could not write temp report file:\n{e}", parent=parent)
            return
        try:
            if sys.platform.startswith("win"):
                os.startfile(path, "print")
            elif sys.platform == "darwin":
                subprocess.Popen(["open", "-a", "TextEdit", path])
            else:
                subprocess.Popen(["xdg-open", path])
        except Exception as e:
            messagebox.showerror(
                "Print failed",
                f"Could not open the print dialog automatically:\n{e}\n\n"
                f"The report was saved here so you can open it manually:\n{path}",
                parent=parent)

    # ---------- Admin auth ----------

    def prompt_admin_login(self, parent=None, reason="This action requires admin authorization."):
        if self.admin_credentials is not None:
            return True
        dialog = tk.Toplevel(parent or self)
        dialog.title("Admin Authorization")
        dialog.resizable(False, False)
        dialog.transient(parent or self)
        dialog.grab_set()
        result = {"ok": False}
        ttk.Label(dialog, text=reason, wraplength=320,
                  justify="left").pack(anchor="w", padx=14, pady=(14, 6))
        ttk.Label(dialog, text="Admin Username:").pack(anchor="w", padx=14, pady=(6, 2))
        user_var = tk.StringVar()
        user_entry = ttk.Entry(dialog, textvariable=user_var, width=30)
        user_entry.pack(padx=14)
        user_entry.focus_set()
        ttk.Label(dialog, text="Admin Password:").pack(anchor="w", padx=14, pady=(8, 2))
        pass_var = tk.StringVar()
        pass_entry = ttk.Entry(dialog, textvariable=pass_var, show="*", width=30)
        pass_entry.pack(padx=14)
        status_var = tk.StringVar(value="")
        ttk.Label(dialog, textvariable=status_var,
                  foreground="#a00000").pack(anchor="w", padx=14, pady=(6, 0))
        button_row = ttk.Frame(dialog)
        button_row.pack(fill="x", padx=14, pady=14)

        def do_cancel():
            result["ok"] = False
            dialog.destroy()

        def do_login():
            username = user_var.get().strip()
            password = pass_var.get().strip()
            if not username or not password:
                status_var.set("Enter both username and password")
                return

            def worker():
                try:
                    r = requests.post(f"{SERVER}/login",
                                      json={"username": username, "password": password},
                                      timeout=5)
                except Exception as e:
                    self.after(0, lambda: status_var.set(f"Server not reachable: {e}"))
                    return

                def apply():
                    if r.status_code == 200:
                        self.admin_credentials = (username, password)
                        result["ok"] = True
                        dialog.destroy()
                    else:
                        try:
                            msg = r.json().get("message", r.text)
                        except Exception:
                            msg = r.text
                        status_var.set(msg or "Invalid credentials")
                self.after(0, apply)
            threading.Thread(target=worker, daemon=True).start()

        ttk.Button(button_row, text="Cancel", command=do_cancel).pack(side="right")
        ttk.Button(button_row, text="Authorize", command=do_login).pack(side="right", padx=(0, 8))
        dialog.bind("<Return>", lambda e: do_login())
        dialog.bind("<Escape>", lambda e: do_cancel())
        center_window_on_screen(dialog, 380, 270)
        dialog.after(150, lambda w=dialog: apply_green_titlebar(w))
        dialog.wait_window()
        return result["ok"]

    def logout_admin(self):
        self.admin_credentials = None

    # ---------- Ribbons ----------

    def show_inventory_ribbon(self):
        self.clear_sub_ribbon()
        ttk.Button(self.sub_ribbon, text="Stocks", image=self.button_icon_stock,
                   compound="top", command=self.build_stocks_content).pack(side="left", padx=3)
        ttk.Button(self.sub_ribbon, text="Products", image=self.button_icon_products,
                   compound="top", command=self.build_products_content).pack(side="left", padx=3)
        ttk.Button(self.sub_ribbon, text="Labels", image=self.button_icon_labels,
                   compound="top", command=self.build_labels_content).pack(side="left", padx=3)
        ttk.Button(self.sub_ribbon, text="Shelfs", image=self.button_icon_shelfs,
                   compound="top", command=self.build_shelfs_content).pack(side="left", padx=3)
        ttk.Button(self.sub_ribbon, text="Admin Panel",
                   compound="top", command=self.open_admin_panel).pack(side="left", padx=3)
        self.switch_group("inventory")

    def open_admin_panel(self):
        """Launches admin_panel.py as a separate process."""
        panel_path = os.path.join(APP_DIR, "admin_panel.py")
        if not os.path.exists(panel_path):
            messagebox.showerror("Missing file",
                                 f"admin_panel.py not found at:\n{panel_path}")
            return
        try:
            subprocess.Popen([sys.executable, panel_path])
        except Exception as e:
            messagebox.showerror("Error", f"Could not launch admin panel:\n{e}")

    def show_purchase_ribbon(self):
        self.clear_sub_ribbon()
        ttk.Button(self.sub_ribbon, text="GRNs", image=self.button_icon_grn,
                   compound="top", command=self.build_grns_content).pack(side="left", padx=3)
        ttk.Button(self.sub_ribbon, text="CRNs", image=self.button_icon_crn,
                   compound="top", command=self.build_crns_content).pack(side="left", padx=3)
        ttk.Button(self.sub_ribbon, text="New Item", image=self.button_icon_products,
                   compound="top", command=self.build_new_item_content).pack(side="left", padx=3)
        self.switch_group("purchase")

    def show_orders_ribbon(self):
        self.clear_sub_ribbon()
        self.switch_group("orders")
        self.build_orders_content()

    def show_pos_ribbon(self):
        self.clear_sub_ribbon()
        self.switch_group("pos")
        self.build_pos_content()

    # ---------- Stocks ----------

    def build_stocks_content(self):
        win = self.open_content_window("stocks", "Stocks", group="inventory", icon=self.icon_stock)
        if win is None:
            return
        top_button_frame = ttk.Frame(win)
        top_button_frame.pack(fill="x", padx=10, pady=(10, 0))
        ttk.Button(top_button_frame, text="Stock List",
                   command=self.open_stock_list_window).pack(side="left")
        filter_frame = ttk.Frame(win)
        filter_frame.pack(fill="x", padx=10, pady=10)
        ttk.Label(filter_frame, text="Code/Barcode:").pack(side="left", padx=(0, 2))
        self.stock_code_var = tk.StringVar()
        ttk.Entry(filter_frame, textvariable=self.stock_code_var, width=12).pack(side="left", padx=(0, 10))
        date_frame = ttk.Frame(filter_frame)
        date_frame.pack(side="right", padx=(10, 0))
        ttk.Label(date_frame, text="From:").pack(side="left", padx=(0, 2))
        self.stock_from_var = tk.StringVar()
        self.from_date_picker = DateEntry(
            date_frame, width=11, date_pattern="yyyy-mm-dd",
            textvariable=self.stock_from_var, showweeknumbers=False)
        self.from_date_picker.pack(side="left", padx=(0, 10))
        self.from_date_picker.bind("<<DateEntrySelected>>", self.on_stock_filter_change)
        ttk.Label(date_frame, text="To:").pack(side="left", padx=(0, 2))
        self.stock_to_var = tk.StringVar()
        self.to_date_picker = DateEntry(
            date_frame, width=11, date_pattern="yyyy-mm-dd",
            textvariable=self.stock_to_var, showweeknumbers=False)
        self.to_date_picker.pack(side="left")
        self.to_date_picker.bind("<<DateEntrySelected>>", self.on_stock_filter_change)
        today = datetime.date.today()
        yesterday = today - datetime.timedelta(days=1)
        self.from_date_picker.set_date(yesterday)
        self.to_date_picker.set_date(today)
        ttk.Label(filter_frame, text="Desc:").pack(side="left", padx=(0, 2))
        self.matched_desc_var = tk.StringVar(value="-")
        ttk.Label(filter_frame, textvariable=self.matched_desc_var,
                  relief="sunken", background="white", foreground="#006600",
                  anchor="w").pack(side="left", fill="x", expand=True, padx=(0, 10))
        self.stock_code_var.trace_add("write", self.on_stock_filter_change)
        history_frame = ttk.LabelFrame(win, text="Stock Movement History")
        history_frame.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        self.stock_history_tree = ttk.Treeview(
            history_frame, columns=("date", "type", "in", "out", "balance"),
            show="headings", style="Dark.Treeview")
        self.stock_history_tree.heading("date", text="Date")
        self.stock_history_tree.heading("type", text="Type")
        self.stock_history_tree.heading("in", text="Stock In")
        self.stock_history_tree.heading("out", text="Stock Out")
        self.stock_history_tree.heading("balance", text="Balance")
        self.stock_history_tree.pack(fill="both", expand=True, padx=5, pady=5)
        self.make_columns_responsive(win, self.stock_history_tree, [2, 1, 1, 1, 1])
        balance_frame = ttk.Frame(win)
        balance_frame.pack(fill="x", padx=10, pady=(0, 10))
        self.live_balance_var = tk.StringVar(value="Bal: -")
        bal_button = tk.Label(
            balance_frame, textvariable=self.live_balance_var,
            font=("Arial", 12, "bold"), bg="#e8a0a0", fg="black",
            relief="raised", bd=2, padx=14, pady=4)
        bal_button.pack(side="right")

    def on_stock_filter_change(self, *args):
        if self._stock_code_normalizing:
            return
        self.update_live_balance()
        if self._stock_filter_after_id is not None:
            try:
                self.after_cancel(self._stock_filter_after_id)
            except (ValueError, tk.TclError):
                pass
            self._stock_filter_after_id = None
        code = self.stock_code_var.get().strip()
        if not code:
            if hasattr(self, "stock_history_tree") and self.stock_history_tree.winfo_exists():
                for row in self.stock_history_tree.get_children():
                    self.stock_history_tree.delete(row)
            return
        self._stock_filter_after_id = self.after(300, self.search_stock_history)

    def update_live_balance(self, *args):
        if not hasattr(self, "live_balance_var"):
            return
        try:
            code = self.stock_code_var.get().strip()
        except tk.TclError:
            return
        if not code:
            self.live_balance_var.set("Bal: -")
            self.matched_desc_var.set("-")
            return
        match = self.find_product_by_code(code)
        if match:
            short_code = self.get_short_code(match)
            if short_code and short_code != code:
                self._stock_code_normalizing = True
                self.stock_code_var.set(short_code)
                self._stock_code_normalizing = False
            self.live_balance_var.set(f"Bal: {match['stock']}")
            self.matched_desc_var.set(match["name"])
        else:
            self.live_balance_var.set("Bal: not found")
            self.matched_desc_var.set("-")

    def search_stock_history(self):
        self._stock_filter_after_id = None
        if not hasattr(self, "stock_history_tree") or not self.stock_history_tree.winfo_exists():
            return
        code = self.stock_code_var.get().strip()
        if not code:
            for row in self.stock_history_tree.get_children():
                self.stock_history_tree.delete(row)
            return
        params = {"barcode": code}
        from_date = self.stock_from_var.get().strip()
        to_date = self.stock_to_var.get().strip()
        if from_date:
            params["from_date"] = from_date
        if to_date:
            params["to_date"] = to_date

        def worker():
            try:
                r = requests.get(f"{SERVER}/stock_ledger", params=params, timeout=3)
                rows = r.json()
            except Exception as e:
                print("Could not reach server:", e)
                return
            self.after(0, lambda: self._apply_stock_history(rows))
        threading.Thread(target=worker, daemon=True).start()

    def _apply_stock_history(self, rows):
        if not hasattr(self, "stock_history_tree") or not self.stock_history_tree.winfo_exists():
            return
        for row in self.stock_history_tree.get_children():
            self.stock_history_tree.delete(row)
        for row in rows:
            self.stock_history_tree.insert(
                "", "end",
                values=(row["date"], row["transaction_type"],
                        row["stock_in"], row["stock_out"], row["balance"]))

    def open_stock_list_window(self):
        win = self.open_content_window("stock_list", "Stock List",
                                       group="inventory", icon=self.icon_stock)
        if win is None:
            return
        center_window_on_screen(win, 380, 200)
        win.minsize(340, 180)
        win.resizable(False, False)
        body = ttk.Frame(win)
        body.pack(fill="both", expand=True, padx=20, pady=20)
        ttk.Button(body, text="Departmental Summary", width=24,
                   command=self.open_departmental_summary_report).pack(pady=(0, 8))
        ttk.Button(body, text="Departments", width=24,
                   command=self.open_departments_picker).pack(pady=(0, 8))
        ttk.Button(body, text="Suppliers", width=24,
                   command=self.open_suppliers_picker).pack()

    # ---------- Departmental Summary ----------

    def open_departmental_summary_report(self):
        def worker():
            try:
                r = requests.get(f"{SERVER}/products", timeout=5)
                data = r.json()
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Error", f"Could not reach server: {e}"))
                return

            def apply():
                self.products_cache = data
                self._show_product_report(data, f"Departmental Summary — All items ({len(data)})")
            self.after(0, apply)
        threading.Thread(target=worker, daemon=True).start()

    def _show_product_report(self, products, title):
        columns = [
            ("code", "Short Code"),
            ("name", "Description"),
            ("unit", "Unit of Measure"),
            ("stock", "Current Stock"),
        ]
        rows = []
        for p in products:
            rows.append((
                self.get_short_code(p),
                p.get("name", ""),
                p.get("packaging", "") or "-",
                p.get("stock", 0),
            ))
        self._open_report_window(
            title=title,
            rows=rows,
            columns=columns,
            weights=[1, 4, 2, 1],
            scope_label=title,
        )

    # ---------- Departments picker ----------

    def open_departments_picker(self):
        win = tk.Toplevel(self)
        win.title("Departments")
        win.transient(self)
        center_window_on_screen(win, 460, 400)
        win.minsize(400, 320)
        win.after(150, lambda w=win: apply_green_titlebar(w))
        top = ttk.Frame(win)
        top.pack(fill="x", padx=12, pady=(12, 0))
        ttk.Label(top, text="Select a department, then click OK.").pack(side="left")
        list_frame = ttk.Frame(win)
        list_frame.pack(fill="both", expand=True, padx=12, pady=10)
        tree = ttk.Treeview(list_frame, columns=("name",), show="headings")
        tree.heading("name", text="Department")
        tree.pack(fill="both", expand=True, side="left")
        vsb = ttk.Scrollbar(list_frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")

        def load_departments():
            for row in tree.get_children():
                tree.delete(row)

            def worker():
                try:
                    r = requests.get(f"{SERVER}/departments", timeout=3)
                    depts = r.json()
                except Exception as e:
                    depts = []
                    print("Could not fetch departments:", e)
                names = {d["name"] if isinstance(d, dict) else str(d) for d in depts}
                for p in self.products_cache:
                    c = str(p.get("category", "")).strip()
                    if c:
                        names.add(c)

                def apply():
                    if not tree.winfo_exists():
                        return
                    for n in sorted(names):
                        tree.insert("", "end", values=(n,))
                self.after(0, apply)
            threading.Thread(target=worker, daemon=True).start()

        load_departments()
        button_row = ttk.Frame(win)
        button_row.pack(fill="x", padx=12, pady=(0, 12))
        ttk.Button(button_row, text="Add Department",
                   command=lambda: self.open_add_department_dialog(
                       parent=win, on_added=load_departments)).pack(side="left")
        ttk.Button(button_row, text="Cancel", command=win.destroy).pack(side="right")

        def do_ok():
            sel = tree.selection()
            if not sel:
                messagebox.showerror("Error", "Select a department first", parent=win)
                return
            name = tree.item(sel[0], "values")[0]
            win.destroy()
            self.open_department_report(name)
        ttk.Button(button_row, text="OK", command=do_ok).pack(side="right", padx=(0, 8))
        tree.bind("<Double-1>", lambda e: do_ok())

    def open_department_report(self, department):
        items = [p for p in self.products_cache
                 if str(p.get("category", "")).strip() == department]
        title = f"Department: {department} — {len(items)} item(s)"
        self._show_product_report(items, title)

    # ---------- Suppliers picker ----------

    def open_suppliers_picker(self):
        win = tk.Toplevel(self)
        win.title("Suppliers")
        win.transient(self)
        center_window_on_screen(win, 460, 400)
        win.minsize(400, 320)
        win.after(150, lambda w=win: apply_green_titlebar(w))
        top = ttk.Frame(win)
        top.pack(fill="x", padx=12, pady=(12, 0))
        ttk.Label(top, text="Select a supplier, then click OK.").pack(side="left")
        list_frame = ttk.Frame(win)
        list_frame.pack(fill="both", expand=True, padx=12, pady=10)
        tree = ttk.Treeview(list_frame, columns=("name",), show="headings")
        tree.heading("name", text="Supplier")
        tree.pack(fill="both", expand=True, side="left")
        vsb = ttk.Scrollbar(list_frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")

        def load_suppliers():
            for row in tree.get_children():
                tree.delete(row)

            def worker():
                try:
                    r = requests.get(f"{SERVER}/suppliers", timeout=3)
                    sups = r.json()
                except Exception as e:
                    sups = []
                    print("Could not fetch suppliers:", e)
                names = {s["name"] if isinstance(s, dict) else str(s) for s in sups}
                for p in self.products_cache:
                    s = str(p.get("supplier", "")).strip()
                    if s:
                        names.add(s)

                def apply():
                    if not tree.winfo_exists():
                        return
                    for n in sorted(names):
                        tree.insert("", "end", values=(n,))
                self.after(0, apply)
            threading.Thread(target=worker, daemon=True).start()

        load_suppliers()
        button_row = ttk.Frame(win)
        button_row.pack(fill="x", padx=12, pady=(0, 12))
        ttk.Button(button_row, text="Add Supplier",
                   command=lambda: self.open_add_supplier_dialog(on_added=load_suppliers)).pack(side="left")
        ttk.Button(button_row, text="Cancel", command=win.destroy).pack(side="right")

        def do_ok():
            sel = tree.selection()
            if not sel:
                messagebox.showerror("Error", "Select a supplier first", parent=win)
                return
            name = tree.item(sel[0], "values")[0]
            win.destroy()
            self.open_supplier_report(name)
        ttk.Button(button_row, text="OK", command=do_ok).pack(side="right", padx=(0, 8))
        tree.bind("<Double-1>", lambda e: do_ok())

    def open_supplier_report(self, supplier):
        items = [p for p in self.products_cache
                 if str(p.get("supplier", "")).strip() == supplier]
        title = f"Supplier: {supplier} — {len(items)} item(s)"
        self._show_product_report(items, title)

    # ---------- Products ----------

    def build_products_content(self):
        win = self.open_content_window("products", "Products", group="inventory", icon=self.icon_products)
        if win is None:
            return
        header_frame = ttk.Frame(win)
        header_frame.pack(fill="x", padx=10, pady=(10, 10))
        code_desc_frame = ttk.LabelFrame(header_frame, text="")
        code_desc_frame.pack(side="left", fill="x", expand=True)
        ttk.Label(code_desc_frame, text="Code/Barcode:").pack(side="left", padx=(5, 2), pady=5)
        self.product_code_var = tk.StringVar()
        ttk.Entry(code_desc_frame, textvariable=self.product_code_var, width=14).pack(side="left", padx=(0, 10), pady=5)
        ttk.Label(code_desc_frame, text="Description:").pack(side="left", padx=(0, 2))
        self.product_desc_var = tk.StringVar(value="-")
        ttk.Entry(code_desc_frame, textvariable=self.product_desc_var, width=28,
                  state="readonly").pack(side="left", fill="x", expand=True, padx=(0, 5), pady=5)
        department_frame = ttk.LabelFrame(header_frame, text="Department")
        department_frame.pack(side="left", padx=(15, 0))
        self.product_department_var = tk.StringVar(value="-")
        ttk.Entry(department_frame, textvariable=self.product_department_var, width=16,
                  state="readonly").pack(padx=8, pady=5)
        packaging_frame = ttk.Frame(win)
        packaging_frame.pack(fill="x", padx=10)
        ttk.Label(packaging_frame, text="Packaging:").pack(side="right", padx=(5, 5))
        self.product_packaging_var = tk.StringVar(value="-")
        ttk.Entry(packaging_frame, textvariable=self.product_packaging_var, width=16,
                  state="readonly").pack(side="right")
        pricing_frame = ttk.LabelFrame(win, text="Pricing (fed via GRN - read only here)")
        pricing_frame.pack(fill="x", padx=10, pady=10)
        ttk.Label(pricing_frame, text="Cost Price:").grid(row=0, column=0, sticky="w", padx=5, pady=5)
        self.product_cost_var = tk.StringVar(value="-")
        ttk.Entry(pricing_frame, textvariable=self.product_cost_var, width=12,
                  state="readonly").grid(row=0, column=1, sticky="w", padx=5, pady=5)
        ttk.Label(pricing_frame, text="Markup:").grid(row=1, column=0, sticky="w", padx=5, pady=5)
        self.product_markup_var = tk.StringVar(value="-")
        ttk.Entry(pricing_frame, textvariable=self.product_markup_var, width=12,
                  state="readonly").grid(row=1, column=1, sticky="w", padx=5, pady=5)
        ttk.Label(pricing_frame, text="Selling Price:").grid(row=2, column=0, sticky="w", padx=5, pady=5)
        self.product_selling_var = tk.StringVar(value="-")
        ttk.Entry(pricing_frame, textvariable=self.product_selling_var, width=12,
                  state="readonly").grid(row=2, column=1, sticky="w", padx=5, pady=5)
        button_frame = ttk.Frame(win)
        button_frame.pack(fill="x", padx=10, pady=(0, 5))
        ttk.Button(button_frame, text="Supplier",
                   command=self.show_product_supplier).pack(side="left", padx=3)
        ttk.Button(button_frame, text="Shelf ID",
                   command=self.show_product_shelf).pack(side="left", padx=3)
        ttk.Button(button_frame, text="Button A").pack(side="left", padx=3)
        ttk.Button(button_frame, text="Button B").pack(side="left", padx=3)
        ttk.Button(button_frame, text="Button C").pack(side="left", padx=3)
        bal_row = ttk.Frame(win)
        bal_row.pack(fill="x", padx=10, pady=(5, 5))
        self.product_live_balance_var = tk.StringVar(value="Bal: -")
        self.product_bal_label = tk.Label(
            bal_row, textvariable=self.product_live_balance_var,
            font=("Arial", 12, "bold"), bg="#e8a0a0", fg="black",
            relief="raised", bd=2, padx=14, pady=4, cursor="hand2")
        self.product_bal_label.pack(side="left")
        self.product_lookup_result_var = tk.StringVar(value="")
        ttk.Label(win, textvariable=self.product_lookup_result_var,
                  foreground="#006600").pack(fill="x", padx=10, pady=(0, 10))
        self.product_code_var.trace_add("write", self.on_product_code_change)

    def on_product_code_change(self, *args):
        if self._product_code_normalizing:
            return
        code = self.product_code_var.get().strip()
        self.product_lookup_result_var.set("")
        if not code:
            self.product_desc_var.set("-")
            self.product_department_var.set("-")
            self.product_packaging_var.set("-")
            self.product_cost_var.set("-")
            self.product_markup_var.set("-")
            self.product_selling_var.set("-")
            self.product_live_balance_var.set("Bal: -")
            return
        match = self.find_product_by_code(code)
        if match:
            short_code = self.get_short_code(match)
            if short_code and short_code != code:
                self._product_code_normalizing = True
                self.product_code_var.set(short_code)
                self._product_code_normalizing = False
            self.product_desc_var.set(match.get("name", "-"))
            self.product_department_var.set(match.get("category", "-"))
            self.product_packaging_var.set(match.get("packaging", "-"))
            self.product_cost_var.set(str(match.get("cost_price", match.get("price", "-"))))
            self.product_markup_var.set(str(match.get("markup", "-")))
            self.product_selling_var.set(str(match.get("selling_price", "-")))
            self.product_live_balance_var.set(f"Bal: {match.get('stock', 0)}")
        else:
            self.product_desc_var.set("Not found")
            self.product_department_var.set("-")
            self.product_packaging_var.set("-")
            self.product_cost_var.set("-")
            self.product_markup_var.set("-")
            self.product_selling_var.set("-")
            self.product_live_balance_var.set("Bal: not found")

    def show_product_supplier(self):
        code = self.product_code_var.get().strip()
        if not code:
            self.product_lookup_result_var.set("Enter a code first")
            return
        match = self.find_product_by_code(code)
        if not match:
            self.product_lookup_result_var.set(f"No product found for code '{code}'")
            return
        supplier = match.get("supplier", "Not set")
        self.product_lookup_result_var.set(f"Supplier: {supplier}")

    def show_product_shelf(self):
        code = self.product_code_var.get().strip()
        if not code:
            self.product_lookup_result_var.set("Enter a code first")
            return
        match = self.find_product_by_code(code)
        if not match:
            self.product_lookup_result_var.set(f"No product found for code '{code}'")
            return
        shelf = match.get("shelf", "Not set")
        self.product_lookup_result_var.set(f"Shelf ID: {shelf}")

    def open_product_list_window(self):
        win = self.open_content_window("product_list", "Product List", group="inventory", icon=self.icon_products)
        if win is None:
            return
        win.minsize(500, 350)
        filter_frame = ttk.Frame(win)
        filter_frame.pack(fill="x", padx=10, pady=10)
        ttk.Label(filter_frame, text="Quick Find (code or name):").pack(side="left")
        filter_var = tk.StringVar()
        ttk.Entry(filter_frame, textvariable=filter_var, width=20).pack(side="left", padx=5)
        tree = ttk.Treeview(win, columns=("id", "name", "cat", "price", "stock", "barcode"), show="headings")
        for col in ("id", "name", "cat", "price", "stock", "barcode"):
            tree.heading(col, text=col.capitalize())
        tree.pack(fill="both", expand=True, padx=10, pady=10)
        self.make_columns_responsive(win, tree, [1, 3, 2, 1, 1, 2])
        after_id = {"id": None}

        def do_refresh():
            after_id["id"] = None
            query = filter_var.get().strip().lower()
            if not tree.winfo_exists():
                return
            for row in tree.get_children():
                tree.delete(row)
            if not query:
                return

            def worker():
                try:
                    r = requests.get(f"{SERVER}/products", timeout=3)
                    products = r.json()
                except Exception as e:
                    self.after(0, lambda: messagebox.showerror("Error", f"Could not reach server: {e}"))
                    return

                def apply():
                    if not tree.winfo_exists():
                        return
                    for p in products:
                        if query in p["barcode"].lower() or query in p["name"].lower():
                            tree.insert("", "end", values=(p["id"], p["name"], p["category"], p["price"], p["stock"], p["barcode"]))
                self.after(0, apply)
            threading.Thread(target=worker, daemon=True).start()

        def on_filter_change(*args):
            if after_id["id"] is not None:
                try:
                    win.after_cancel(after_id["id"])
                except (ValueError, tk.TclError):
                    pass
            after_id["id"] = win.after(300, do_refresh)
        filter_var.trace_add("write", on_filter_change)

    def open_products_by_supplier_window(self):
        win = self.open_content_window("products_by_supplier", "Products by Supplier",
                                       group="inventory", icon=self.icon_products)
        if win is None:
            return
        win.minsize(500, 350)
        filter_frame = ttk.Frame(win)
        filter_frame.pack(fill="x", padx=10, pady=10)
        ttk.Label(filter_frame, text="Supplier:").pack(side="left", padx=(0, 5))
        supplier_var = tk.StringVar()
        supplier_combo = ttk.Combobox(filter_frame, textvariable=supplier_var, width=30, state="readonly")
        supplier_combo.pack(side="left")
        tree = ttk.Treeview(win, columns=("code", "name", "cat", "price", "stock", "barcode"), show="headings")
        for col, label in zip(
            ("code", "name", "cat", "price", "stock", "barcode"),
            ("Code", "Name", "Category", "Price", "Stock", "Barcode")):
            tree.heading(col, text=label)
        tree.pack(fill="both", expand=True, padx=10, pady=10)
        self.make_columns_responsive(win, tree, [1, 3, 2, 1, 1, 2])

        def refresh_list():
            if not tree.winfo_exists():
                return
            for row in tree.get_children():
                tree.delete(row)
            selected = supplier_var.get().strip()
            if not selected:
                return
            for p in self.products_cache:
                if str(p.get("supplier", "")).strip() == selected:
                    tree.insert("", "end", values=(
                        self.get_short_code(p), p.get("name", ""), p.get("category", ""),
                        p.get("price", p.get("selling_price", "")),
                        p.get("stock", ""), p.get("barcode", "")))
        ttk.Button(filter_frame, text="OK", command=refresh_list).pack(side="left", padx=(8, 0))

        def worker():
            try:
                r = requests.get(f"{SERVER}/suppliers", timeout=3)
                suppliers = r.json()
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Error", f"Could not reach server: {e}"))
                return

            def apply():
                if not supplier_combo.winfo_exists():
                    return
                names = [s["name"] if isinstance(s, dict) else s for s in suppliers]
                supplier_combo["values"] = names
            self.after(0, apply)
        threading.Thread(target=worker, daemon=True).start()

    def open_products_by_department_window(self):
        win = self.open_content_window("products_by_department", "Products by Department",
                                       group="inventory", icon=self.icon_products)
        if win is None:
            return
        win.minsize(500, 350)
        filter_frame = ttk.Frame(win)
        filter_frame.pack(fill="x", padx=10, pady=10)
        ttk.Label(filter_frame, text="Department:").pack(side="left", padx=(0, 5))
        dept_var = tk.StringVar()
        dept_combo = ttk.Combobox(filter_frame, textvariable=dept_var, width=30, state="readonly")
        dept_combo.pack(side="left")
        tree = ttk.Treeview(win, columns=("code", "name", "cat", "price", "stock", "barcode"), show="headings")
        for col, label in zip(
            ("code", "name", "cat", "price", "stock", "barcode"),
            ("Code", "Name", "Category", "Price", "Stock", "Barcode")):
            tree.heading(col, text=label)
        tree.pack(fill="both", expand=True, padx=10, pady=10)
        self.make_columns_responsive(win, tree, [1, 3, 2, 1, 1, 2])

        def refresh_list():
            if not tree.winfo_exists():
                return
            for row in tree.get_children():
                tree.delete(row)
            selected = dept_var.get().strip()
            if not selected:
                return
            for p in self.products_cache:
                if str(p.get("category", "")).strip() == selected:
                    tree.insert("", "end", values=(
                        self.get_short_code(p), p.get("name", ""), p.get("category", ""),
                        p.get("price", p.get("selling_price", "")),
                        p.get("stock", ""), p.get("barcode", "")))
        ttk.Button(filter_frame, text="OK", command=refresh_list).pack(side="left", padx=(8, 0))
        ttk.Button(filter_frame, text="+ New Department",
                   command=lambda: self.open_add_department_dialog(
                       parent=win, target_var=dept_var,
                       on_added=lambda: self._refresh_department_list(dept_combo))).pack(side="left", padx=(8, 0))
        self._refresh_department_list(dept_combo)

    def _refresh_department_list(self, combo):
        if not combo.winfo_exists():
            return
        names = {str(p.get("category", "")).strip()
                 for p in self.products_cache if str(p.get("category", "")).strip()}

        def worker():
            try:
                r = requests.get(f"{SERVER}/departments", timeout=3)
                depts = r.json()
            except Exception as e:
                print("Could not fetch departments:", e)
                depts = []

            def apply():
                if not combo.winfo_exists():
                    return
                for d in depts:
                    n = d["name"] if isinstance(d, dict) else str(d)
                    if n.strip():
                        names.add(n.strip())
                current = combo.get()
                combo["values"] = sorted(names)
                if current in names:
                    combo.set(current)
            self.after(0, apply)
        threading.Thread(target=worker, daemon=True).start()

    def open_add_department_dialog(self, parent=None, target_var=None, on_added=None):
        dialog = tk.Toplevel(parent or self)
        dialog.title("New Department")
        dialog.resizable(False, False)
        dialog.transient(parent or self)
        dialog.grab_set()
        center_window_on_screen(dialog, 340, 180)
        dialog.after(150, lambda w=dialog: apply_green_titlebar(w))
        ttk.Label(dialog, text="Department Name:").pack(anchor="w", padx=12, pady=(15, 2))
        name_var = tk.StringVar()
        name_entry = ttk.Entry(dialog, textvariable=name_var, width=32)
        name_entry.pack(padx=12)
        name_entry.focus_set()
        status_var = tk.StringVar(value="")
        ttk.Label(dialog, textvariable=status_var, foreground="#a00000").pack(anchor="w", padx=12, pady=(6, 0))
        button_row = ttk.Frame(dialog)
        button_row.pack(fill="x", padx=12, pady=15)

        def do_save():
            name = name_var.get().strip()
            if not name:
                status_var.set("Enter a department name")
                return

            def worker():
                try:
                    r = requests.post(f"{SERVER}/departments", json={"name": name}, timeout=5)
                except Exception as e:
                    self.after(0, lambda: status_var.set(f"Server not reachable: {e}"))
                    return

                def apply():
                    if r.status_code in (200, 201) or r.status_code == 409:
                        self.known_departments.add(name)
                        if target_var is not None:
                            target_var.set(name)
                        if on_added:
                            on_added()
                        dialog.destroy()
                    else:
                        status_var.set(r.text or "Could not save department")
                self.after(0, apply)
            threading.Thread(target=worker, daemon=True).start()
        ttk.Button(button_row, text="Cancel", command=dialog.destroy).pack(side="right")
        ttk.Button(button_row, text="Save", command=do_save).pack(side="right", padx=(0, 8))

    # =======================================================================
    # LABELS WINDOW
    # =======================================================================

    def build_labels_content(self):
        win = self.open_content_window("labels", "Labels", group="inventory", icon=self.icon_labels)
        if win is None:
            return

        self._label_win = win
        self.label_mode_var = tk.StringVar(value="Product Label")
        self.label_mode_buttons = {}
        self._label_code_normalizing = False
        self._label_bottom_built = False
        self._label_current_area = None

        try:
            win.minsize(720, 360)
            center_window_on_screen(win, 760, 400)
        except Exception:
            pass

        mode_row = ttk.Frame(win)
        mode_row.pack(fill="x", padx=10, pady=(10, 6), side="top")

        def make_mode_button(label):
            btn = ttk.Button(mode_row, text=label,
                             command=lambda l=label: self.set_label_mode(l))
            btn.pack(side="left", padx=(0, 4))
            self.label_mode_buttons[label] = btn

        make_mode_button("Product Label")
        make_mode_button("Shelf Label")
        make_mode_button("Test Barcode")

        self._label_body = ttk.Frame(win)
        self._label_body.pack(fill="both", expand=True, padx=0, pady=0, side="top")

        btn_row = ttk.Frame(win)
        btn_row.pack(fill="x", padx=10, pady=(6, 10), side="bottom")
        ttk.Button(btn_row, text="Change (select printer)",
                   command=self.change_label_printer).pack(side="left")
        ttk.Button(btn_row, text="Print",
                   command=self.print_labels).pack(side="right", ipadx=24, ipady=4)
        self._label_bottom_built = True

        self.set_label_mode("Product Label")

    def _build_label_print_area(self):
        if self._label_current_area == "label":
            return
        for w in self._label_body.winfo_children():
            w.destroy()

        row_code = ttk.Frame(self._label_body)
        row_code.pack(fill="x", padx=10, pady=(4, 4))

        ttk.Label(row_code, text="Code:", font=("Arial", 10, "bold")).pack(side="left", padx=(0, 6))
        self.label_barcode_var = tk.StringVar()
        self.label_barcode_entry = ttk.Entry(row_code, textvariable=self.label_barcode_var, width=22)
        self.label_barcode_entry.pack(side="left", padx=(0, 16))
        self.label_barcode_entry.focus_set()

        ttk.Label(row_code, text="Description:", font=("Arial", 10, "bold")).pack(side="left", padx=(0, 6))
        self.label_desc_var = tk.StringVar(value="")
        ttk.Entry(row_code, textvariable=self.label_desc_var, width=42).pack(side="left", fill="x", expand=True)

        row_main = ttk.Frame(self._label_body)
        row_main.pack(fill="both", expand=True, padx=10, pady=(4, 4))

        qty_col = ttk.Frame(row_main)
        qty_col.pack(side="left", anchor="n", padx=(0, 16))
        ttk.Label(qty_col, text="Quantity of Labels to print:",
                  font=("Arial", 10, "bold")).pack(anchor="w", pady=(0, 4))
        self.label_qty_var = tk.StringVar(value="1")
        ttk.Entry(qty_col, textvariable=self.label_qty_var, width=8).pack(anchor="w")

        preview_col = ttk.LabelFrame(row_main, text="Label preview")
        preview_col.pack(side="left", fill="both", expand=True)

        self.label_preview_canvas = tk.Canvas(
            preview_col, width=380, height=190,
            bg="white", highlightthickness=0)
        self.label_preview_canvas.pack(padx=8, pady=6)

        self.label_barcode_var.trace_add("write", self.on_label_code_change)
        self.label_desc_var.trace_add("write", lambda *a: self._refresh_label_preview())
        self.label_mode_var.trace_add("write", lambda *a: self._refresh_label_preview())

        self._refresh_label_preview()
        self._label_current_area = "label"

    def _build_label_test_area(self):
        if self._label_current_area == "test":
            return
        for w in self._label_body.winfo_children():
            w.destroy()

        row_code = ttk.Frame(self._label_body)
        row_code.pack(fill="x", padx=10, pady=(4, 6))

        ttk.Label(row_code, text="Code:", font=("Arial", 10, "bold")).pack(side="left", padx=(0, 6))
        self.label_barcode_var = tk.StringVar()
        self.label_barcode_entry = ttk.Entry(row_code, textvariable=self.label_barcode_var, width=22)
        self.label_barcode_entry.pack(side="left")
        self.label_barcode_entry.focus_set()

        table_frame = ttk.LabelFrame(self._label_body, text="Scanned item")
        table_frame.pack(fill="both", expand=True, padx=10, pady=(4, 8))

        self.label_display_tree = ttk.Treeview(
            table_frame,
            columns=("code", "desc", "cost", "selling"),
            show="headings",
            height=8,
            style="Dark.Treeview")
        self.label_display_tree.heading("code", text="Code")
        self.label_display_tree.heading("desc", text="Description")
        self.label_display_tree.heading("cost", text="Cost")
        self.label_display_tree.heading("selling", text="Selling price")
        self.label_display_tree.column("code", width=130, anchor="w")
        self.label_display_tree.column("desc", width=320, anchor="w")
        self.label_display_tree.column("cost", width=110, anchor="e")
        self.label_display_tree.column("selling", width=130, anchor="e")
        self.label_display_tree.pack(fill="both", expand=True, padx=5, pady=5)

        self.label_barcode_var.trace_add("write", self.on_label_code_change)
        self._label_current_area = "test"

    def set_label_mode(self, mode):
        self.label_mode_var.set(mode)
        for label, btn in self.label_mode_buttons.items():
            if label == mode:
                btn.state(["pressed"])
            else:
                btn.state(["!pressed"])
        if mode == "Test Barcode":
            self._build_label_test_area()
        else:
            self._build_label_print_area()

    def on_label_code_change(self, *args):
        if self._label_code_normalizing:
            return
        code = self.label_barcode_var.get().strip()
        mode = self.label_mode_var.get() if self.label_mode_var else "Product Label"

        if mode == "Test Barcode":
            if hasattr(self, "label_display_tree") and self.label_display_tree.winfo_exists():
                for row in self.label_display_tree.get_children():
                    self.label_display_tree.delete(row)
            if not code:
                return
            match = self.find_product_by_code(code)
            if not match:
                self.label_display_tree.insert("", "end",
                                               values=("Not found", "-", "-", "-"))
                return
            short_code = self.get_short_code(match)
            desc = match.get("name", "")
            cost = match.get("cost_price", match.get("price", ""))
            selling = match.get("selling_price", match.get("price", ""))

            def fmt(v):
                try:
                    return f"{float(v):.2f}"
                except (TypeError, ValueError):
                    return str(v)
            self.label_display_tree.insert(
                "", "end", values=(short_code, desc, fmt(cost), fmt(selling)))
        else:
            if not code:
                if hasattr(self, "label_desc_var"):
                    self.label_desc_var.set("")
                self._refresh_label_preview()
                return
            match = self.find_product_by_code(code)
            if hasattr(self, "label_desc_var"):
                self.label_desc_var.set(match.get("name", "") if match else "")
            self._refresh_label_preview()

    # -------- on-screen preview --------

    def _refresh_label_preview(self):
        if not hasattr(self, "label_preview_canvas") or not self.label_preview_canvas.winfo_exists():
            return
        c = self.label_preview_canvas
        c.delete("all")

        W = 380
        H = 190
        mode = self.label_mode_var.get() if self.label_mode_var else "Product Label"
        code = self.label_barcode_var.get().strip() if hasattr(self, "label_barcode_var") else ""
        desc = self.label_desc_var.get().strip() if hasattr(self, "label_desc_var") else ""
        product = self.find_product_by_code(code) if code else None

        if mode == "Product Label":
            gap = 8
            label_w = (W - 16 - gap) // 2
            label_h = label_w // 2
            if label_h > H - 16:
                label_h = H - 16
                label_w = label_h * 2
            self._draw_preview_product(c, 8, 8, label_w, label_h, desc, product, code)
            self._draw_preview_product(c, 8 + label_w + gap, 8, label_w, label_h, desc, product, code)
        elif mode == "Shelf Label":
            label_w = 260
            label_h = label_w // 2
            x = (W - label_w) // 2
            y = (H - label_h) // 2
            self._draw_preview_shelf(c, x, y, label_w, label_h, desc, product, code)

        c.create_rectangle(2, 2, W - 2, H - 2, outline="#999999")

    def _draw_preview_product(self, c, x0, y0, w, h, desc, product, code):
        pad = 4
        cx = x0 + w / 2

        banner_h = int(h * 0.20)
        desc_h = int(h * 0.16)
        bc_h = int(h * 0.30)
        code_h = int(h * 0.14)

        y = y0 + pad

        c.create_rectangle(x0 + pad, y, x0 + w - pad, y + banner_h,
                           fill="black", outline="black")
        c.create_text(cx, y + banner_h / 2, text=SHOP_NAME,
                      font=("Arial", 9, "bold"), fill="white")
        y += banner_h + 1

        c.create_text(cx, y, text=(desc or "(description)")[:22],
                      font=("Arial", 7, "bold"), fill="black", anchor="n")
        y += desc_h

        short_code = ""
        selling = ""
        if product:
            short_code = self.get_short_code(product)
            selling = product.get("selling_price", product.get("price", ""))

        long_code = code or "0000000000000"
        bits = code128_encode(long_code)
        bc_x0 = x0 + pad + 4
        bc_x1 = x0 + w - pad - 4
        n = max(1, len(bits))
        bw = (bc_x1 - bc_x0) / n
        for i, b in enumerate(bits):
            if b == "1":
                bx = bc_x0 + i * bw
                c.create_rectangle(bx, y, bx + max(bw, 0.4), y + bc_h,
                                   fill="black", outline="")
        y += bc_h + 1

        if short_code:
            c.create_text(cx, y, text=short_code,
                          font=("Courier", 7, "bold"), fill="black", anchor="n")
        y += code_h

        try:
            selling_str = f"KSh {float(selling):.2f}"
        except (TypeError, ValueError):
            selling_str = f"KSh {selling}" if selling else "KSh -"
        c.create_text(cx, y0 + h - pad - 2, text=selling_str,
                      font=("Arial", 11, "bold"), fill="#006600", anchor="s")

    def _draw_preview_shelf(self, c, x0, y0, w, h, desc, product, code):
        pad = 6
        cx = x0 + w / 2

        banner_h = int(h * 0.18)
        desc_h = int(h * 0.16)
        bc_h = int(h * 0.28)
        code_h = int(h * 0.14)

        y = y0 + pad
        c.create_rectangle(x0 + pad, y, x0 + w - pad, y + banner_h,
                           fill="black", outline="black")
        c.create_text(cx, y + banner_h / 2, text=SHOP_NAME,
                      font=("Arial", 11, "bold"), fill="white")
        y += banner_h + 2

        c.create_text(cx, y, text=(desc or "(description)")[:30],
                      font=("Arial", 9, "bold"), fill="black", anchor="n")
        y += desc_h

        short_code = ""
        selling = ""
        if product:
            short_code = self.get_short_code(product)
            selling = product.get("selling_price", product.get("price", ""))

        long_code = code or "0000000000000"
        bits = code128_encode(long_code)
        bc_x0 = x0 + pad + 12
        bc_x1 = x0 + w - pad - 12
        n = max(1, len(bits))
        bw = (bc_x1 - bc_x0) / n
        for i, b in enumerate(bits):
            if b == "1":
                bx = bc_x0 + i * bw
                c.create_rectangle(bx, y, bx + max(bw, 0.4), y + bc_h,
                                   fill="black", outline="")
        y += bc_h + 2

        if short_code:
            c.create_text(cx, y, text=f"Code: {short_code}",
                          font=("Courier", 10, "bold"), fill="black", anchor="n")
        y += code_h

        try:
            selling_str = f"KSh {float(selling):.2f}"
        except (TypeError, ValueError):
            selling_str = f"KSh {selling}" if selling else "KSh -"
        c.create_text(cx, y0 + h - pad - 2, text=selling_str,
                      font=("Arial", 14, "bold"), fill="#006600", anchor="s")

    # -------- printer + PDF --------

    def change_label_printer(self):
        if not sys.platform.startswith("win"):
            messagebox.showinfo("Change printer",
                                "Printer selection uses the OS print dialog when "
                                "you click Print on this platform.")
            return
        try:
            from ctypes import wintypes

            class PRINTDLGW(ctypes.Structure):
                _fields_ = [
                    ("lStructSize", wintypes.DWORD),
                    ("hwndOwner", wintypes.HWND),
                    ("hDevMode", wintypes.HANDLE),
                    ("hDevNames", wintypes.HANDLE),
                    ("hDC", wintypes.HDC),
                    ("Flags", wintypes.DWORD),
                    ("nFromPage", wintypes.WORD),
                    ("nToPage", wintypes.WORD),
                    ("nMinPage", wintypes.WORD),
                    ("nMaxPage", wintypes.WORD),
                    ("nCopies", wintypes.WORD),
                    ("hInstance", wintypes.HANDLE),
                    ("lCustData", ctypes.c_void_p),
                    ("lpfnPrintHook", ctypes.c_void_p),
                    ("lpfnSetupHook", ctypes.c_void_p),
                    ("lpPrintTemplateName", wintypes.LPCWSTR),
                    ("lpSetupTemplateName", wintypes.LPCWSTR),
                    ("hPrintTemplate", wintypes.HANDLE),
                    ("hSetupTemplate", wintypes.HANDLE),
                ]

            PD_RETURNDC = 0x00000100
            PD_USEDEVMODECOPIESANDCOLLATE = 0x00040000
            PD_NOSELECTION = 0x00000004
            PD_NOPAGENUMS = 0x00000008

            self.update_idletasks()
            hwnd_int = int(self.winfo_id())
            real_hwnd = ctypes.windll.user32.GetParent(hwnd_int)
            if not real_hwnd:
                real_hwnd = hwnd_int

            pd = PRINTDLGW()
            pd.lStructSize = ctypes.sizeof(PRINTDLGW)
            pd.hwndOwner = real_hwnd
            pd.Flags = (PD_RETURNDC
                        | PD_USEDEVMODECOPIESANDCOLLATE
                        | PD_NOSELECTION
                        | PD_NOPAGENUMS)
            pd.nMinPage = 1
            pd.nMaxPage = 1

            comdlg32 = ctypes.windll.comdlg32
            comdlg32.PrintDlgW.argtypes = [ctypes.POINTER(PRINTDLGW)]
            comdlg32.PrintDlgW.restype = wintypes.BOOL

            ok = comdlg32.PrintDlgW(ctypes.byref(pd))
            if not ok:
                return

            if pd.hDevNames:
                kernel32 = ctypes.windll.kernel32
                kernel32.GlobalLock.argtypes = [wintypes.HGLOBAL]
                kernel32.GlobalLock.restype = ctypes.c_void_p
                kernel32.GlobalUnlock.argtypes = [wintypes.HGLOBAL]
                kernel32.GlobalUnlock.restype = wintypes.BOOL

                ptr = kernel32.GlobalLock(pd.hDevNames)
                try:
                    if ptr:
                        offsets = ctypes.cast(
                            ctypes.c_void_p(ptr),
                            ctypes.POINTER(ctypes.c_ushort)
                        )
                        device_offset = offsets[1]
                        device_addr = ptr + device_offset
                        device_name = ctypes.cast(
                            ctypes.c_void_p(device_addr),
                            ctypes.c_wchar_p
                        ).value
                        if device_name:
                            self.label_printer_name = device_name
                finally:
                    kernel32.GlobalUnlock(pd.hDevNames)

            if pd.hDC:
                ctypes.windll.gdi32.DeleteDC(pd.hDC)

            messagebox.showinfo(
                "Printer selected",
                f"Selected printer:\n{self.label_printer_name or '(unknown)'}")

        except Exception as e:
            messagebox.showwarning(
                "Couldn't open printer dialog",
                f"{e}\n\nYou can still print - the Save dialog will let you "
                f"pick a target file and print from your PDF viewer.")

    def _label_sizes(self):
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        return {
            "pagesize": A4,
            "label_w": 50 * mm,
            "label_h": 25 * mm,
            "cols": 2,
            "rows": 7,
            "margin_x": 8 * mm,
            "margin_y": 10 * mm,
            "per_page": 2 * 7,
        }

    def _draw_barcode(self, canvas, bits, x, y, width, height):
        n = len(bits)
        if n == 0:
            return
        bw = width / n
        canvas.saveState()
        canvas.setFillColorRGB(0, 0, 0)
        for i, b in enumerate(bits):
            if b == "1":
                canvas.rect(x + i * bw, y, bw + 0.05, height, stroke=0, fill=1)
        canvas.restoreState()

    def _draw_product_label(self, canvas, x, y, size, product):
        from reportlab.lib import colors

        label_w = size["label_w"]
        label_h = size["label_h"]
        cx = x + label_w / 2

        canvas.saveState()

        banner_h = label_h * 0.22
        y_top = y + label_h

        canvas.setFillColor(colors.black)
        canvas.rect(x, y_top - banner_h, label_w, banner_h, stroke=0, fill=1)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 9)
        canvas.drawCentredString(cx, y_top - banner_h + 5, SHOP_NAME)

        canvas.setFillColor(colors.black)
        canvas.setFont("Helvetica-Bold", 7)
        desc = (product.get("name") or "")[:32]
        desc_y = y_top - banner_h - 8
        canvas.drawCentredString(cx, desc_y, desc)

        bc_top = desc_y - 4
        bc_h = label_h * 0.32
        bc_bottom = bc_top - bc_h
        long_code = str(product.get("barcode") or product.get("code") or "")
        bits = code128_encode(long_code)
        bc_x = x + 4
        bc_w = label_w - 8
        n = max(1, len(bits))
        bw = bc_w / n
        canvas.setFillColor(colors.black)
        for i, b in enumerate(bits):
            if b == "1":
                canvas.rect(bc_x + i * bw, bc_bottom, bw + 0.05, bc_h, stroke=0, fill=1)

        canvas.setFillColor(colors.black)
        canvas.setFont("Courier-Bold", 7)
        short = str(product.get("code") or product.get("barcode") or "")
        code_y = bc_bottom - 8
        canvas.drawCentredString(cx, code_y, short)

        selling = product.get("selling_price", product.get("price", ""))
        try:
            selling_str = f"KSh {float(selling):.2f}"
        except (TypeError, ValueError):
            selling_str = f"KSh {selling}"
        canvas.setFillColor(colors.HexColor("#006600"))
        canvas.setFont("Helvetica-Bold", 12)
        canvas.drawCentredString(cx, y + 4, selling_str)

        canvas.restoreState()

    def _draw_shelf_label(self, canvas, x, y, size, product):
        from reportlab.lib import colors

        label_w = size["label_w"]
        label_h = size["label_h"]
        cx = x + label_w / 2

        canvas.saveState()

        banner_h = label_h * 0.22
        y_top = y + label_h

        canvas.setFillColor(colors.black)
        canvas.rect(x, y_top - banner_h, label_w, banner_h, stroke=0, fill=1)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 10)
        canvas.drawCentredString(cx, y_top - banner_h + 5, SHOP_NAME)

        canvas.setFillColor(colors.black)
        canvas.setFont("Helvetica-Bold", 8)
        desc = (product.get("name") or "")[:28]
        desc_y = y_top - banner_h - 9
        canvas.drawCentredString(cx, desc_y, desc)

        bc_top = desc_y - 4
        bc_h = label_h * 0.32
        bc_bottom = bc_top - bc_h
        long_code = str(product.get("barcode") or product.get("code") or "")
        bits = code128_encode(long_code)
        bc_x = x + 6
        bc_w = label_w - 12
        n = max(1, len(bits))
        bw = bc_w / n
        canvas.setFillColor(colors.black)
        for i, b in enumerate(bits):
            if b == "1":
                canvas.rect(bc_x + i * bw, bc_bottom, bw + 0.05, bc_h, stroke=0, fill=1)

        canvas.setFont("Courier-Bold", 7)
        short = str(product.get("code") or product.get("barcode") or "")
        code_y = bc_bottom - 8
        canvas.drawCentredString(cx, code_y, f"Code: {short}")

        selling = product.get("selling_price", product.get("price", ""))
        try:
            selling_str = f"KSh {float(selling):.2f}"
        except (TypeError, ValueError):
            selling_str = f"KSh {selling}"
        canvas.setFillColor(colors.HexColor("#006600"))
        canvas.setFont("Helvetica-Bold", 13)
        canvas.drawCentredString(cx, y + 4, selling_str)

        canvas.restoreState()

    def print_labels(self):
        mode = self.label_mode_var.get() if self.label_mode_var else "Product Label"

        if mode == "Test Barcode":
            messagebox.showinfo(
                "Test Barcode",
                "Test Barcode mode is display-only - it shows the info "
                "resolved from the long code on screen without printing.")
            return

        code = self.label_barcode_var.get().strip()
        if not code:
            messagebox.showerror("Error", "Scan or type the long barcode first")
            return

        product = self.find_product_by_code(code)
        if not product:
            messagebox.showerror("Error", f"No product found for code '{code}'")
            return

        try:
            qty = int(self.label_qty_var.get())
            if qty <= 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("Error", "Quantity must be a positive whole number")
            return

        try:
            from reportlab.pdfgen import canvas as rl_canvas
            from reportlab.lib.units import mm
        except ImportError:
            messagebox.showerror("Missing library",
                                 "Labels need 'reportlab'.\n\n"
                                 "Install it with:   pip install reportlab")
            return

        default_name = f"labels_{mode.replace(' ', '_').lower()}_{code}.pdf"
        path = filedialog.asksaveasfilename(
            parent=self,
            title=f"Save {mode} PDF",
            defaultextension=".pdf",
            initialfile=default_name,
            filetypes=[("PDF Document", "*.pdf"), ("All files", "*.*")])
        if not path:
            return

        size = self._label_sizes()
        c = rl_canvas.Canvas(path, pagesize=size["pagesize"])
        page_w, page_h = size["pagesize"]
        per_page = size["per_page"]
        origin_x = size["margin_x"]
        origin_y = page_h - size["margin_y"] - size["label_h"]

        if mode == "Product Label":
            total_rows = (qty + 1) // 2
            for r in range(total_rows):
                for col in range(2):
                    if r * 2 + col >= qty:
                        break
                    x = origin_x + col * size["label_w"]
                    y = origin_y - r * size["label_h"]
                    self._draw_product_label(c, x, y, size, product)
                    c.saveState()
                    c.setStrokeColorRGB(0.85, 0.85, 0.85)
                    c.setLineWidth(0.3)
                    c.rect(x, y, size["label_w"], size["label_h"], stroke=1, fill=0)
                    c.restoreState()

                if (r + 1) % size["rows"] == 0 and r < total_rows - 1:
                    c.showPage()
                    origin_y = page_h - size["margin_y"] - size["label_h"]
        else:
            for i in range(qty):
                row = i % size["rows"]
                y = origin_y - row * size["label_h"]
                x = origin_x
                self._draw_shelf_label(c, x, y, size, product)
                c.saveState()
                c.setStrokeColorRGB(0.85, 0.85, 0.85)
                c.setLineWidth(0.3)
                c.rect(x, y, size["label_w"], size["label_h"], stroke=1, fill=0)
                c.restoreState()

                if (i + 1) % size["rows"] == 0 and i < qty - 1:
                    c.showPage()
                    origin_y = page_h - size["margin_y"] - size["label_h"]

        c.save()

        messagebox.showinfo("Labels ready", f"Saved {qty} label(s) to:\n{path}")
        open_with_default_app(path)

    def build_shelfs_content(self):
        win = self.open_content_window("shelfs", "Shelfs", group="inventory", icon=self.icon_shelfs)
        if win is None:
            return
        ttk.Label(win, text="Shelfs — coming soon", font=("Arial", 14)).pack(pady=40)

    # ---------- Purchase > GRNs ----------

    def build_grns_content(self):
        win = self.open_content_window("grns", "GRNs", group="purchase", icon=self.icon_grn)
        if win is None:
            return
        self.grn_lines = []
        self.grn_supplier_locked = False
        supplier_frame = ttk.Frame(win)
        supplier_frame.pack(fill="x", padx=10, pady=(10, 0))
        ttk.Button(supplier_frame, text="GRN List", command=self.open_grn_list_window).pack(side="right")
        ttk.Label(supplier_frame, text="Select Supplier:").pack(side="left", padx=(0, 5))
        self.grn_supplier_var = tk.StringVar()
        self.grn_supplier_combo = ttk.Combobox(supplier_frame, textvariable=self.grn_supplier_var, width=30)
        self.grn_supplier_combo.pack(side="left")
        self.fetch_grn_suppliers()
        self.grn_new_supplier_btn = ttk.Button(
            supplier_frame, text="+ New Supplier",
            command=lambda: self.open_add_supplier_dialog(
                target_var=self.grn_supplier_var, on_added=self.fetch_grn_suppliers))
        self.grn_new_supplier_btn.pack(side="left", padx=(8, 0))
        self.grn_supplier_lock_btn = ttk.Button(
            supplier_frame, text="Lock Supplier", command=self.toggle_grn_supplier_lock)
        self.grn_supplier_lock_btn.pack(side="left", padx=(8, 0))
        self.grn_locked_supplier_var = tk.StringVar(value="")
        ttk.Label(supplier_frame, textvariable=self.grn_locked_supplier_var,
                  foreground="#006600", font=("Arial", 9, "bold")).pack(side="left", padx=(10, 0))
        line_frame = ttk.Frame(win)
        line_frame.pack(fill="x", padx=10, pady=10)
        qty_pack_frame = ttk.Frame(line_frame)
        qty_pack_frame.pack(side="right")
        ttk.Label(qty_pack_frame, text="Quantity:").grid(row=0, column=0, sticky="w", padx=(0, 2))
        ttk.Label(qty_pack_frame, text="Packaging:").grid(row=0, column=1, sticky="w", padx=(10, 2))
        self.grn_qty_var = tk.StringVar()
        self.grn_packaging_var = tk.StringVar()
        ttk.Entry(qty_pack_frame, textvariable=self.grn_qty_var, width=8).grid(row=1, column=0, sticky="w")
        self.grn_packaging_combo = ttk.Combobox(qty_pack_frame, textvariable=self.grn_packaging_var, width=16)
        self.grn_packaging_combo.grid(row=1, column=1, sticky="w", padx=(10, 0))
        code_col = ttk.Frame(line_frame)
        code_col.pack(side="left", padx=(0, 10))
        ttk.Label(code_col, text="Code / Barcode:").pack(anchor="w")
        self.grn_code_var = tk.StringVar()
        ttk.Entry(code_col, textvariable=self.grn_code_var, width=12).pack()
        desc_col = ttk.Frame(line_frame)
        desc_col.pack(side="left", fill="x", expand=True, padx=(0, 10))
        ttk.Label(desc_col, text="Description:").pack(anchor="w")
        self.grn_desc_var = tk.StringVar(value="-")
        ttk.Entry(desc_col, textvariable=self.grn_desc_var, state="readonly").pack(fill="x", expand=True)
        pricing_row = ttk.Frame(win)
        pricing_row.pack(fill="x", padx=10, pady=(0, 10))
        ttk.Button(pricing_row, text="Delete", command=self.remove_grn_line).pack(side="right", padx=(5, 8))
        ttk.Button(pricing_row, text="Add", command=self.add_grn_line).pack(side="right", pady=8)
        self.grn_code_var.trace_add("write", self.on_grn_code_change)
        bottom_frame = ttk.Frame(win)
        bottom_frame.pack(side="bottom", fill="x", padx=10, pady=10)
        self.grn_total_var = tk.StringVar(value="Value of all delivery: 0.00")
        total_label = tk.Label(bottom_frame, textvariable=self.grn_total_var,
                               font=("Arial", 12, "bold"), bg="#e8a0a0", fg="black",
                               relief="raised", bd=2, padx=14, pady=4)
        total_label.pack(side="left")
        receive_button = tk.Label(bottom_frame, text="RECEIVE", font=("Arial", 11, "bold"),
                                  bg="#8fd08f", fg="black", relief="raised", bd=2, padx=16, pady=6,
                                  cursor="hand2")
        receive_button.pack(side="right")
        receive_button.bind("<Button-1>", lambda e: self.receive_grn())
        ttk.Button(bottom_frame, text="Clear", command=self.clear_grn_form).pack(side="right", padx=(0, 10))
        lines_frame = ttk.LabelFrame(win, text="Items fed but not yet added to system stock")
        lines_frame.pack(side="top", fill="both", expand=True, padx=10, pady=(0, 5))
        tree_container = ttk.Frame(lines_frame)
        tree_container.pack(fill="both", expand=True, padx=5, pady=5)
        self.grn_tree = ttk.Treeview(
            tree_container, columns=("code", "desc", "qty", "cost", "total"),
            show="headings", height=6)
        self.grn_tree.heading("code", text="Code")
        self.grn_tree.heading("desc", text="Description")
        self.grn_tree.heading("qty", text="Quantity")
        self.grn_tree.heading("cost", text="Cost")
        self.grn_tree.heading("total", text="Total")
        grn_scrollbar = ttk.Scrollbar(tree_container, orient="vertical", command=self.grn_tree.yview)
        self.grn_tree.configure(yscrollcommand=grn_scrollbar.set)
        self.grn_tree.pack(side="left", fill="both", expand=True)
        grn_scrollbar.pack(side="right", fill="y")
        self.make_columns_responsive(win, self.grn_tree, [2, 4, 1, 1, 1])

    def fetch_grn_suppliers(self):
        def worker():
            try:
                r = requests.get(f"{SERVER}/suppliers", timeout=3)
                suppliers = r.json()
            except Exception as e:
                print("Could not fetch suppliers:", e)
                return
            self.after(0, lambda: self._apply_grn_suppliers(suppliers))
        threading.Thread(target=worker, daemon=True).start()

    def _apply_grn_suppliers(self, suppliers):
        if not hasattr(self, "grn_supplier_combo") or not self.grn_supplier_combo.winfo_exists():
            return
        names = [s["name"] if isinstance(s, dict) else s for s in suppliers]
        self.grn_supplier_combo["values"] = names

    def toggle_grn_supplier_lock(self):
        if not self.grn_supplier_locked:
            supplier = self.grn_supplier_var.get().strip()
            if not supplier:
                messagebox.showerror("Error", "Select a supplier first")
                return
            self.grn_supplier_locked = True
            self.grn_supplier_combo.configure(state="disabled")
            self.grn_new_supplier_btn.configure(state="disabled")
            self.grn_supplier_lock_btn.configure(text="Change Supplier")
            self.grn_locked_supplier_var.set(f"Locked to: {supplier}")
        else:
            self._reset_grn_supplier_lock()

    def _reset_grn_supplier_lock(self):
        self.grn_supplier_locked = False
        if hasattr(self, "grn_supplier_combo") and self.grn_supplier_combo.winfo_exists():
            self.grn_supplier_combo.configure(state="normal")
        if hasattr(self, "grn_new_supplier_btn") and self.grn_new_supplier_btn.winfo_exists():
            self.grn_new_supplier_btn.configure(state="normal")
        if hasattr(self, "grn_supplier_lock_btn") and self.grn_supplier_lock_btn.winfo_exists():
            self.grn_supplier_lock_btn.configure(text="Lock Supplier")
        if hasattr(self, "grn_locked_supplier_var"):
            self.grn_locked_supplier_var.set("")

    def open_add_supplier_dialog(self, target_var=None, on_added=None):
        if target_var is None and hasattr(self, "grn_supplier_var"):
            target_var = self.grn_supplier_var
        if on_added is None and hasattr(self, "fetch_grn_suppliers"):
            on_added = self.fetch_grn_suppliers
        dialog = tk.Toplevel(self)
        dialog.title("Add New Supplier")
        dialog.resizable(False, False)
        dialog.transient(self)
        dialog.grab_set()
        center_window_on_screen(dialog, 320, 180)
        dialog.after(150, lambda w=dialog: apply_green_titlebar(w))
        ttk.Label(dialog, text="Supplier Name:").pack(anchor="w", padx=12, pady=(15, 2))
        name_var = tk.StringVar()
        name_entry = ttk.Entry(dialog, textvariable=name_var, width=32)
        name_entry.pack(padx=12)
        name_entry.focus_set()
        ttk.Label(dialog, text="Phone (optional):").pack(anchor="w", padx=12, pady=(10, 2))
        phone_var = tk.StringVar()
        ttk.Entry(dialog, textvariable=phone_var, width=32).pack(padx=12)
        button_row = ttk.Frame(dialog)
        button_row.pack(fill="x", padx=12, pady=15)
        ttk.Button(button_row, text="Cancel", command=dialog.destroy).pack(side="right")
        ttk.Button(button_row, text="Save",
                   command=lambda: self.submit_new_supplier(dialog, name_var, phone_var, target_var, on_added)).pack(side="right", padx=(0, 8))

    def submit_new_supplier(self, dialog, name_var, phone_var, target_var=None, on_added=None):
        name = name_var.get().strip()
        if not name:
            messagebox.showerror("Error", "Enter a supplier name", parent=dialog)
            return
        payload = {"name": name}
        phone = phone_var.get().strip()
        if phone:
            payload["phone"] = phone

        def worker():
            try:
                r = requests.post(f"{SERVER}/suppliers", json=payload, timeout=5)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Error", f"Could not reach server: {e}", parent=dialog))
                return

            def apply():
                if r.status_code in (200, 201):
                    if on_added:
                        on_added()
                    if target_var is not None:
                        target_var.set(name)
                    dialog.destroy()
                else:
                    messagebox.showerror("Error", r.text, parent=dialog)
            self.after(0, apply)
        threading.Thread(target=worker, daemon=True).start()

    def on_grn_code_change(self, *args):
        if self._grn_code_normalizing:
            return
        code = self.grn_code_var.get().strip()
        if not code:
            self.grn_desc_var.set("-")
            return
        match = self.find_product_by_code(code)
        if match:
            short_code = self.get_short_code(match)
            if short_code and short_code != code:
                self._grn_code_normalizing = True
                self.grn_code_var.set(short_code)
                self._grn_code_normalizing = False
            self.grn_desc_var.set(match["name"])
            packaging_options = []
            existing_packaging = match.get("packaging")
            if existing_packaging:
                packaging_options.append(existing_packaging)
            for generic in ("Loose", "Carton", "Box", "Pack"):
                if generic not in packaging_options:
                    packaging_options.append(generic)
            self.grn_packaging_combo["values"] = packaging_options
            if existing_packaging and not self.grn_packaging_var.get().strip():
                self.grn_packaging_var.set(existing_packaging)
        else:
            self.grn_desc_var.set("Not found — add it via Purchase > New Item")

    def add_grn_line(self):
        if not self.grn_supplier_locked:
            messagebox.showerror("Error", "Select a supplier and click 'Lock Supplier' before adding items")
            return
        code = self.grn_code_var.get().strip()
        if not code:
            messagebox.showerror("Error", "Enter a product code")
            return
        match = self.find_product_by_code(code)
        if not match:
            messagebox.showerror("Error", f"No product found for code '{code}'")
            return
        try:
            qty = float(self.grn_qty_var.get())
            if qty <= 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("Error", "Enter a valid quantity greater than 0")
            return

        def to_float(value, default=0.0):
            try:
                return float(value)
            except (TypeError, ValueError):
                return default
        unit_cost = to_float(match.get("cost_price", match.get("price")))
        markup = to_float(match.get("markup"))
        markup_type = match.get("markup_type", "%")
        selling_price = to_float(match.get("selling_price", match.get("price")))
        packaging = self.grn_packaging_var.get().strip()
        line_total = qty * unit_cost
        self.grn_lines.append({
            "barcode": match["barcode"],
            "name": match["name"],
            "qty": qty,
            "packaging": packaging,
            "unit_cost": unit_cost,
            "markup": markup,
            "markup_type": markup_type,
            "selling_price": selling_price,
            "line_total": line_total,
        })
        self.grn_tree.insert("", "end",
                             values=(match["barcode"], match["name"], qty,
                                     f"{unit_cost:.2f}", f"{line_total:.2f}"))
        self.recompute_grn_total()
        self.grn_code_var.set("")
        self.grn_qty_var.set("")
        self.grn_packaging_var.set("")
        self.grn_desc_var.set("-")

    def remove_grn_line(self):
        selected = self.grn_tree.selection()
        if not selected:
            return
        for item_id in selected:
            index = self.grn_tree.index(item_id)
            self.grn_tree.delete(item_id)
            if 0 <= index < len(self.grn_lines):
                del self.grn_lines[index]
        self.recompute_grn_total()

    def recompute_grn_total(self):
        total = sum(line["line_total"] for line in self.grn_lines)
        self.grn_total_var.set(f"Value of all delivery: {total:.2f}")

    def clear_grn_form(self):
        self.grn_lines = []
        if hasattr(self, "grn_tree") and self.grn_tree.winfo_exists():
            for row in self.grn_tree.get_children():
                self.grn_tree.delete(row)
        self.grn_code_var.set("")
        self.grn_qty_var.set("")
        self.grn_packaging_var.set("")
        self.grn_desc_var.set("-")
        self._reset_grn_supplier_lock()
        self.grn_supplier_var.set("")
        self.recompute_grn_total()

    def receive_grn(self):
        if not self.grn_lines:
            messagebox.showerror("Error", "Add at least one item before receiving")
            return
        supplier = self.grn_supplier_var.get().strip()
        if not supplier:
            messagebox.showerror("Error", "Select a supplier")
            return
        payload = {
            "supplier": supplier,
            "date": datetime.date.today().isoformat(),
            "lines": [{
                "barcode": line["barcode"], "qty": line["qty"],
                "packaging": line["packaging"], "unit_cost": line["unit_cost"],
                "markup": line["markup"], "markup_type": line["markup_type"],
                "selling_price": line["selling_price"],
            } for line in self.grn_lines],
        }

        def worker():
            try:
                r = requests.post(f"{SERVER}/grn", json=payload, timeout=5)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Error", f"Could not reach server: {e}"))
                return

            def apply():
                if r.status_code in (200, 201):
                    messagebox.showinfo("Success", "Stock received and prices updated")
                    self.refresh_products_cache()
                    self.clear_grn_form()
                else:
                    messagebox.showerror("Error", r.text)
            self.after(0, apply)
        threading.Thread(target=worker, daemon=True).start()

    def open_grn_list_window(self):
        win = self.open_content_window("grn_list", "GRN List", group="purchase", icon=self.icon_grn)
        if win is None:
            return
        win.minsize(400, 300)
        tree = ttk.Treeview(win, columns=("id", "supplier", "date", "total"), show="headings")
        tree.heading("id", text="GRN #")
        tree.heading("supplier", text="Supplier")
        tree.heading("date", text="Date")
        tree.heading("total", text="Total")
        tree.pack(fill="both", expand=True, padx=10, pady=10)
        self.make_columns_responsive(win, tree, [1, 3, 2, 1])

        def worker():
            try:
                r = requests.get(f"{SERVER}/grns", timeout=3)
                grns = r.json()
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Error", f"Could not reach server: {e}"))
                return

            def apply():
                if not tree.winfo_exists():
                    return
                for g in grns:
                    tree.insert("", "end", values=(g.get("id"), g.get("supplier"), g.get("date"), g.get("total")))
            self.after(0, apply)
        threading.Thread(target=worker, daemon=True).start()

    def build_crns_content(self):
        win = self.open_content_window("crns", "CRNs", group="purchase", icon=self.icon_crn)
        if win is None:
            return
        ttk.Label(win, text="CRNs (Credit/Return Notes) — coming soon", font=("Arial", 14)).pack(pady=40)

    # ---------- Purchase > New Item ----------

    def build_new_item_content(self):
        win = self.open_content_window("new_item", "New Item", group="purchase", icon=self.icon_products)
        if win is None:
            return
        if not self.prompt_admin_login(
                parent=win,
                reason="Adding a new item to the system requires admin authorization."):
            win.destroy()
            self.open_windows.pop("new_item", None)
            self.window_groups.pop("new_item", None)
            return
        self.new_item_status_var = tk.StringVar(value="")
        form = ttk.Frame(win)
        form.pack(fill="both", expand=True, padx=15, pady=15)
        supplier_row = ttk.Frame(form)
        supplier_row.pack(fill="x", pady=(0, 12))
        ttk.Label(supplier_row, text="Supplier:").pack(side="left", padx=(0, 5))
        self.new_item_supplier_var = tk.StringVar()
        self.new_item_supplier_combo = ttk.Combobox(
            supplier_row, textvariable=self.new_item_supplier_var, width=28)
        self.new_item_supplier_combo.pack(side="left")
        ttk.Button(supplier_row, text="+ New Supplier",
                   command=lambda: self.open_add_supplier_dialog(
                       target_var=self.new_item_supplier_var,
                       on_added=self.fetch_new_item_suppliers)).pack(side="left", padx=(8, 0))
        self.fetch_new_item_suppliers()
        fields_frame = ttk.Frame(form)
        fields_frame.pack(fill="x", pady=(0, 10))
        self.new_item_barcode_var = tk.StringVar()
        self.new_item_code_var = tk.StringVar()
        self.new_item_name_var = tk.StringVar()
        self.new_item_category_var = tk.StringVar()
        self.new_item_packaging_var = tk.StringVar()
        self.new_item_cost_var = tk.StringVar()
        self.new_item_markup_var = tk.StringVar()
        self.new_item_markup_type_var = tk.StringVar(value="%")
        self.new_item_selling_var = tk.StringVar(value="-")

        def add_row(row, label_text, var, combo_values=None):
            ttk.Label(fields_frame, text=label_text).grid(row=row, column=0, sticky="w", padx=(0, 8), pady=4)
            if combo_values is not None:
                widget = ttk.Combobox(fields_frame, textvariable=var, values=combo_values, width=26)
            else:
                widget = ttk.Entry(fields_frame, textvariable=var, width=28)
            widget.grid(row=row, column=1, sticky="w", pady=4)
            return widget
        add_row(0, "Barcode (long, scanned):", self.new_item_barcode_var)
        add_row(1, "Short Code (internal, optional):", self.new_item_code_var)
        add_row(2, "Description / Name:", self.new_item_name_var)
        add_row(3, "Department / Category:", self.new_item_category_var)
        add_row(4, "Packaging:", self.new_item_packaging_var, combo_values=["Loose", "Carton", "Box", "Pack"])
        add_row(5, "Initial Cost Price:", self.new_item_cost_var)
        ttk.Label(fields_frame, text="Markup:").grid(row=6, column=0, sticky="w", padx=(0, 8), pady=4)
        markup_input_frame = ttk.Frame(fields_frame)
        markup_input_frame.grid(row=6, column=1, sticky="w", pady=4)
        ttk.Entry(markup_input_frame, textvariable=self.new_item_markup_var, width=10).pack(side="left")
        ttk.Combobox(markup_input_frame, textvariable=self.new_item_markup_type_var,
                     values=["%", "Amount"], width=7, state="readonly").pack(side="left", padx=(6, 0))
        ttk.Label(fields_frame, text="Selling Price:").grid(row=7, column=0, sticky="w", padx=(0, 8), pady=4)
        ttk.Entry(fields_frame, textvariable=self.new_item_selling_var, width=28,
                  state="readonly").grid(row=7, column=1, sticky="w", pady=4)
        self.new_item_cost_var.trace_add("write", self.recompute_new_item_selling_price)
        self.new_item_markup_var.trace_add("write", self.recompute_new_item_selling_price)
        self.new_item_markup_type_var.trace_add("write", self.recompute_new_item_selling_price)
        ttk.Label(form, textvariable=self.new_item_status_var,
                  foreground="#006600", wraplength=680).pack(fill="x", pady=(0, 5))
        button_row = ttk.Frame(form)
        button_row.pack(fill="x")
        ttk.Button(button_row, text="Clear", command=self.clear_new_item_form).pack(side="right", padx=(8, 0))
        ttk.Button(button_row, text="Save Item", command=self.submit_new_item).pack(side="right")
        ttk.Button(button_row, text="Logout", command=self.logout_admin).pack(side="right", padx=(0, 8))

    def fetch_new_item_suppliers(self):
        def worker():
            try:
                r = requests.get(f"{SERVER}/suppliers", timeout=3)
                suppliers = r.json()
            except Exception as e:
                print("Could not fetch suppliers:", e)
                return
            self.after(0, lambda: self._apply_new_item_suppliers(suppliers))
        threading.Thread(target=worker, daemon=True).start()

    def _apply_new_item_suppliers(self, suppliers):
        if not hasattr(self, "new_item_supplier_combo") or not self.new_item_supplier_combo.winfo_exists():
            return
        names = [s["name"] if isinstance(s, dict) else s for s in suppliers]
        self.new_item_supplier_combo["values"] = names

    def recompute_new_item_selling_price(self, *args):
        try:
            cost = float(self.new_item_cost_var.get())
            markup = float(self.new_item_markup_var.get())
        except ValueError:
            self.new_item_selling_var.set("-")
            return
        if self.new_item_markup_type_var.get() == "Amount":
            selling_price = cost + markup
        else:
            selling_price = cost + (cost * markup / 100)
        self.new_item_selling_var.set(f"{selling_price:.2f}")

    def clear_new_item_form(self):
        self.new_item_barcode_var.set("")
        self.new_item_code_var.set("")
        self.new_item_name_var.set("")
        self.new_item_category_var.set("")
        self.new_item_packaging_var.set("")
        self.new_item_cost_var.set("")
        self.new_item_markup_var.set("")
        self.new_item_markup_type_var.set("%")
        self.new_item_selling_var.set("-")

    def submit_new_item(self):
        if not self.admin_credentials:
            if not self.prompt_admin_login(reason="Adding a new item requires admin authorization."):
                return
        supplier = self.new_item_supplier_var.get().strip()
        if not supplier:
            messagebox.showerror("Error", "Select a supplier (or add a new one)")
            return
        barcode = self.new_item_barcode_var.get().strip()
        short_code = self.new_item_code_var.get().strip() or barcode
        name = self.new_item_name_var.get().strip()
        if not barcode or not name:
            messagebox.showerror("Error", "Barcode and Description are required")
            return
        if self.find_product_by_code(barcode) or self.find_product_by_code(short_code):
            messagebox.showerror("Error", f"A product with code or barcode '{barcode}' already exists")
            return
        try:
            cost_price = float(self.new_item_cost_var.get()) if self.new_item_cost_var.get().strip() else 0.0
            if cost_price < 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("Error", "Enter a valid cost price")
            return
        try:
            markup = float(self.new_item_markup_var.get()) if self.new_item_markup_var.get().strip() else 0.0
        except ValueError:
            messagebox.showerror("Error", "Enter a valid markup")
            return
        markup_type = self.new_item_markup_type_var.get()
        if markup_type == "Amount":
            selling_price = cost_price + markup
        else:
            selling_price = cost_price + (cost_price * markup / 100)
        payload = {
            "barcode": barcode, "code": short_code, "name": name,
            "category": self.new_item_category_var.get().strip(),
            "packaging": self.new_item_packaging_var.get().strip(),
            "supplier": supplier, "cost_price": cost_price,
            "markup": markup, "markup_type": markup_type,
            "selling_price": round(selling_price, 2),
            "price": round(selling_price, 2), "stock": 0,
        }

        def worker():
            try:
                u, p = self.admin_credentials
                headers = {"X-Admin-User": u, "X-Admin-Pass": p}
                r = requests.post(f"{SERVER}/products", json=payload,
                                  headers=headers, timeout=5)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Error", f"Could not reach server: {e}"))
                return

            def apply():
                if r.status_code in (200, 201):
                    self.refresh_products_cache()
                    self.new_item_status_var.set(
                        f"Saved '{name}' (code: {short_code}). It's now available in GRN and elsewhere.")
                    self.clear_new_item_form()
                elif r.status_code == 401:
                    messagebox.showerror("Error", "Admin authorization failed. Please log in again.")
                    self.logout_admin()
                else:
                    messagebox.showerror("Error", r.text)
            self.after(0, apply)
        threading.Thread(target=worker, daemon=True).start()

    def build_orders_content(self):
        win = self.open_content_window("orders", "Orders", group="orders", icon=self.icon_orders)
        if win is None:
            return
        ttk.Label(win, text="Orders — coming soon", font=("Arial", 14)).pack(pady=40)

    def build_pos_content(self):
        win = self.open_content_window("pos", "POS - Live Sales Monitor", group="pos", icon=self.icon_pos)
        if win is None:
            return
        self.sales_list = tk.Listbox(win, font=("Consolas", 11))
        self.sales_list.pack(fill="both", expand=True, padx=10, pady=10)
        self.status_label = ttk.Label(win, text="Monitoring for new sales...")
        self.status_label.pack(pady=5)
        self.poll_sales()

    def poll_sales(self):
        if not hasattr(self, "status_label") or not self.status_label.winfo_exists():
            return

        def worker():
            try:
                r = requests.get(f"{SERVER}/sales", timeout=3)
                sales = r.json()
                error = None
            except Exception as e:
                sales = None
                error = e
            self.after(0, lambda: self._apply_sales_poll(sales, error))
        threading.Thread(target=worker, daemon=True).start()
        self.after(3000, self.poll_sales)

    def _apply_sales_poll(self, sales, error):
        if not hasattr(self, "status_label") or not self.status_label.winfo_exists():
            return
        if error is not None:
            self.status_label.config(text=f"Server not reachable: {error}")
            return
        new_sales = [s for s in sales if s["id"] > self.last_sale_id_seen]
        if new_sales:
            for s in reversed(new_sales):
                line = f"Sale #{s['id']}  |  Total: {s['total']}  |  {s['date']}"
                self.sales_list.insert(0, line)
            self.last_sale_id_seen = max(s["id"] for s in sales)
            self.status_label.config(text="Last update: new sale(s) detected!")
        else:
            self.status_label.config(text="Monitoring for new sales... (no changes)")


if __name__ == "__main__":
    BackOffice().mainloop()